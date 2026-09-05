"""QA 校验：字体度量换行模拟 + PPTX 回读验证。

设计理由
--------
1. 度量与渲染同源：PowerPoint 排版时按字体文件里的字形 advance width 决定一行能
   放下多少字。本模块用 FontTools 读取同一份字体（微软雅黑 msyh.ttc），拿到每个
   字符的真实宽度比例（advance / unitsPerEm，与字号线性无关），再乘以字号换算成
   pt，用与渲染器一致的贪心换行规则模拟行数。度量来源与渲染来源同源，"算出来的
   行数"才能逼近 PowerPoint 实际排出的行数——这是溢出检测误报/漏报可控的前提。
   若凭字符数粗估（如 len(text) / N），中文全宽与拉丁半宽混排时误差会立刻失控。
2. error / warning 分级：
   - error 是结构性缺陷（页数不符、形状超出画布边界）：无论在什么机器、什么字体
     环境下打开都会直接呈现为可见错误，必须修复；
   - warning 是排版质量风险（文本可能溢出文本框）：真实表现取决于用户机器的字体
     回退、行高、内边距等运行时因素，本模块的度量只是高置信度近似，所以降级为
     warning 提示人工复核，而不作为硬失败阻断。
3. 字体读取失败时降级为经验宽度估算（CJK 1.0x 字号、拉丁 0.55x、空格 0.33x），
   不抛异常，只把 used_estimate 置 True 让调用方知晓结果置信度下降。
"""

from functools import lru_cache

from pptx import Presentation

EMU_PER_PT = 12700  # 1 pt = 12700 EMU（OOXML 规范）
DEFAULT_FONT_SIZE_PT = 18.0   # run 上取不到字号时的回退值
LINE_HEIGHT_FACTOR = 1.25     # 行高系数（近似 PowerPoint 默认单倍行距）
OVERFLOW_TOLERANCE_PT = 2.0   # 需求高度超出框高 2pt 以内不判溢出
BOUNDARY_TOLERANCE_PT = 0.5   # 形状越界判断容差
MSYH_PATH = "C:/Windows/Fonts/msyh.ttc"


# ---------------------------------------------------------------- 字体度量

@lru_cache(maxsize=1)
def _load_font():
    """加载微软雅黑，返回 (cmap, hmtx, unitsPerEm)；任何失败返回 None。

    lru_cache(maxsize=1) 保证整个模块生命周期只读一次字体文件。
    """
    try:
        from fontTools.ttLib import TTFont
        font = TTFont(MSYH_PATH, fontNumber=0)
        return font.getBestCmap(), font["hmtx"], font["head"].unitsPerEm
    except Exception:
        return None


@lru_cache(maxsize=8192)
def _char_em_width(ch: str):
    """字符在微软雅黑中的宽度比例（advance / unitsPerEm，字号无关的常数）。

    缺字形或字体不可用时返回 None，由调用方走降级估算。
    """
    data = _load_font()
    if data is None:
        return None
    cmap, hmtx, upem = data
    glyph = cmap.get(ord(ch))
    if glyph is None:
        return None
    return hmtx[glyph][0] / upem


def _char_width_pt(ch: str, size_pt: float) -> float:
    """单字符渲染宽度（pt）。真实度量优先，缺失时降级经验估算。"""
    em = _char_em_width(ch)
    if em is not None:
        return em * size_pt
    if ch == " ":
        return 0.33 * size_pt
    if ord(ch) < 128:
        return 0.55 * size_pt
    return 1.0 * size_pt


def font_available() -> bool:
    """微软雅黑是否读取成功（False 表示当前结果来自降级估算）。"""
    return _load_font() is not None


# ---------------------------------------------------------------- 换行模拟

def _tokenize(text: str):
    """把文本切成 (kind, token) 序列，kind 决定断行策略：
    - "word"：连续 ASCII 字母/数字合成一个词，优先在空格处断行；
    - "cjk"：CJK 汉字与中文标点等非 ASCII 字符，逐字可断；
    - "space"：半角空格，可断且换行后行首丢弃。
    ASCII 标点（如 . , -）不并入拉丁词，按逐字可断处理，避免长串无法折行。
    """
    tokens = []
    buf = []
    kind = None

    def flush():
        nonlocal buf, kind
        if buf:
            tokens.append((kind, "".join(buf)))
            buf = []
            kind = None

    for ch in text:
        if ch == " ":
            flush()
            tokens.append(("space", ch))
        elif ch.isascii() and ch.isalnum():
            if kind != "word":
                flush()
                kind = "word"
            buf.append(ch)
        else:
            flush()
            tokens.append(("cjk", ch))
    flush()
    return tokens


def _place_units(units, size_pt, box_width_pt, first_width):
    """把不可再分的单元序列逐个填行（超长词/超宽单字强拆用）。

    first_width 是首个单元前已占的行宽（进入强拆时恒为 0）。
    返回 (总行数, 末行已占宽度)。
    """
    lines, cur = 1, first_width
    for ch in units:
        w = _char_width_pt(ch, size_pt)
        if cur + w <= box_width_pt or cur == 0:
            cur += w
        else:
            lines += 1
            cur = w
    return lines, cur


def _measure_lines_ex(text: str, size_pt: float, box_width_pt: float):
    """贪心换行模拟，返回 (行数, 是否降级估算)。"""
    tokens = _tokenize(text)
    if not tokens:
        return 1, _load_font() is None

    lines, cur = 1, 0.0
    pending_space = 0.0  # 行中累积的空格宽，遇到下一个内容 token 时一并结算
    for kind, tok in tokens:
        if kind == "space":
            pending_space += _char_width_pt(" ", size_pt)
            continue

        if kind == "word":
            width = sum(_char_width_pt(c, size_pt) for c in tok)
            units = tok          # 超长词按字符强拆
        else:  # cjk / 单字符
            width = _char_width_pt(tok, size_pt)
            units = tok          # 单字超行宽时独占强拆

        if cur == 0:
            pending_space = 0.0  # 行首空格丢弃

        if cur + pending_space + width <= box_width_pt or cur + pending_space == 0:
            # 放得下（含行中空格）；或当前行为空（由强拆/换行负责继续填）
            if cur == 0:
                placed, cur = _place_units(units, size_pt, box_width_pt, 0.0)
                lines += placed - 1
            else:
                cur += pending_space + width
            pending_space = 0.0
            continue

        # 当前行放不下 → 换行，行首空格丢弃
        lines += 1
        pending_space = 0.0
        if width > box_width_pt:
            # 超长单元：新行仍放不下，逐字强拆
            placed, cur = _place_units(units, size_pt, box_width_pt, 0.0)
            lines += placed - 1
        else:
            cur = width

    return lines, _load_font() is None


def measure_text_lines(text: str, size_pt: float, box_width_pt: float) -> int:
    """模拟文本在 box_width_pt 宽的文本框内的换行，返回行数。

    规则：CJK 字符与中文标点逐字可断；连续拉丁字母/数字视为一个词、优先在
    空格处断行；超长词（整词宽超过行宽）逐字强拆；半角空格可断且行首丢弃。
    字体读取失败时内部自动降级为经验宽度估算，不抛异常。
    """
    lines, _ = _measure_lines_ex(text, size_pt, box_width_pt)
    return lines


# ---------------------------------------------------------------- PPTX 校验

def _to_pt(length) -> float:
    """EMU Length → pt。"""
    return int(length) / EMU_PER_PT if length is not None else 0.0


def _para_font_size(paragraph) -> float:
    """段落字号：逐 run 取第一个显式设置的值，取不到回退 18pt。"""
    for run in paragraph.runs:
        if run.font.size is not None:
            return run.font.size.pt
    return DEFAULT_FONT_SIZE_PT


def _check_text_overflow(slide_no: int, shape, warnings: list):
    """文本溢出检查：需求行高（行数 x 字号 x 1.25）与框高比较，超出 +2pt 判溢出。"""
    tf = shape.text_frame
    text = tf.text
    if not text or not text.strip():
        return

    inner_w = _to_pt(shape.width) - _to_pt(tf.margin_left) - _to_pt(tf.margin_right)
    inner_h = _to_pt(shape.height) - _to_pt(tf.margin_top) - _to_pt(tf.margin_bottom)
    if inner_w <= 0:
        return  # 异常几何交给边界检查，这里不重复报警

    needed = 0.0
    for para in tf.paragraphs:
        if not para.text:
            continue
        size = _para_font_size(para)
        lines = measure_text_lines(para.text, size, inner_w)
        needed += lines * size * LINE_HEIGHT_FACTOR
        # 段前/段后距必须计入：builder 要点标题段 space_before=20pt 高频使用，
        # 漏加会让"行高勉强够+段距多"的页系统性漏报（审计 M1）
        needed += _to_pt(para.space_before or 0) + _to_pt(para.space_after or 0)
    if needed > inner_h + OVERFLOW_TOLERANCE_PT:
        preview = text.replace("\n", " ")[:20]
        warnings.append(
            f"第{slide_no}页 形状'{shape.name}' 文字可能溢出："
            f"需约 {needed:.1f}pt 行高，框内高 {inner_h:.1f}pt，文本\"{preview}…\""
        )


def check_pptx(path: str, expected_pages: int) -> dict:
    """python-pptx 回读导出文件，校验页数 / 形状边界 / 文本溢出。

    返回 {"errors": [str], "warnings": [str], "used_estimate": bool}；
    used_estimate=True 表示字体度量走了降级估算，warning 置信度下降。
    """
    errors: list[str] = []
    warnings: list[str] = []

    try:
        prs = Presentation(path)
    except Exception as exc:
        return {
            "errors": [f"无法打开 PPTX：{path}（{exc}）"],
            "warnings": [],
            "used_estimate": False,
        }

    used_estimate = not font_available()

    # 页数校验（error 级：结构性不符）
    actual = len(prs.slides)
    if actual != expected_pages:
        errors.append(f"页数不符：期望 {expected_pages} 页，实际 {actual} 页")

    # 逐页校验
    canvas_w = _to_pt(prs.slide_width)
    canvas_h = _to_pt(prs.slide_height)
    for slide_no, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            # 形状越界（error 级）：left/top/width/height 与画布比较，±0.5pt 容差
            left, top = _to_pt(shape.left), _to_pt(shape.top)
            right, bottom = left + _to_pt(shape.width), top + _to_pt(shape.height)
            if (left < -BOUNDARY_TOLERANCE_PT or top < -BOUNDARY_TOLERANCE_PT
                    or right > canvas_w + BOUNDARY_TOLERANCE_PT
                    or bottom > canvas_h + BOUNDARY_TOLERANCE_PT):
                errors.append(
                    f"第{slide_no}页 形状'{shape.name}' 超出画布边界："
                    f"({left:.1f}, {top:.1f})-({right:.1f}, {bottom:.1f})，"
                    f"画布 ({canvas_w:.1f}, {canvas_h:.1f})"
                )
            # 文本溢出（warning 级）
            if shape.has_text_frame:
                _check_text_overflow(slide_no, shape, warnings)

    return {"errors": errors, "warnings": warnings, "used_estimate": used_estimate}
