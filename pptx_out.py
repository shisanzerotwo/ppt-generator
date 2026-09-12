"""`deck.json` → **可编辑** pptx（母版 + 占位符）。

与 `builder.py` 的分工（D11 / R9）
----------------------------------
`builder.py` 是"从 JSON 重建 + 图片版"的老路径，**一字不改**（既有基线依赖它）；
本模块是第二条导出路径：走 PowerPoint 的母版/占位符，导出的文字是**真文本**，
在 PowerPoint 里可选中、可改、不破版。

中文字体为什么要动 XML（契约 §7.4）
-----------------------------------
python-pptx 默认模板的 `theme1.xml` 里 `a:ea typeface=""`（空）、
`<a:font script="Hans" typeface="宋体"/>`。而 `run.font.name` **只写 `a:latin`**；
中文属东亚文字范围、按 `a:ea` 匹配 → 不处理时演示稿里的中文是**宋体**。
所以双管齐下：
  ① 每个写入的 run 补 `a:ea`（必须 append 在 `a:latin` **之后**，OOXML 元素序要求）；
  ② `prs.save()` 后**后处理 zip**，把 theme 的 `a:ea`/Hans 都改成微软雅黑。
② 是必要的——用户在 PowerPoint 里新敲的字走主题字体，只改 run 救不了后续编辑。

IR 装不下的两类形状（**本模块无法往返**，见 IMPL_REPORT 契约缺陷节）
--------------------------------------------------------------------
`ShapeInfo` 只带几何，不带**图片路径**也不带**图表数据**，所以：
- `picture`：无法 `insert_picture(abs_path)`（IR 里没有 abs_path）；
- `chart`：无法 `add_chart(..., chart_data)`（IR 里没有 series/labels/values）。
两者一律跳过并在 `report["degraded"]` 里点名，不静默丢。文本与表格是完整往返的。
"""

import os
import re
import time
import zipfile
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import builder
import hl_layout
import qa
from pptx_io import DeckIR, PptxError

FONT = builder.FONT
EMU_PER_PT = qa.EMU_PER_PT

# 内置 5 版式 → python-pptx 默认模板的真实 layout 下标（契约 §7.2）
_ROLE_LAYOUT = {"cover": 0, "content": 1, "data": 5, "end": 0}

# 兜底安全区（占位符缺失时用）
_SAFE_LEFT, _SAFE_TOP = Inches(0.6), Inches(0.7)
_SAFE_W, _SAFE_H = Inches(12.1), Inches(6.1)

_TITLE_PH = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE,
             PP_PLACEHOLDER.VERTICAL_TITLE)
_BODY_PH = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE,
            PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT)


@dataclass
class TemplateInfo:
    path: str
    layout_names: list[str]
    # layout_idx -> [(ph_idx, ph_type, name), …]
    placeholders: dict = field(default_factory=dict)
    has_theme_cjk: bool = False


# ---------------------------------------------------------------- 模板

def _theme_xml(prs) -> str:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    try:
        part = prs.slide_masters[0].part.part_related_by(RT.THEME)
        return part.blob.decode("utf-8", errors="replace")
    except Exception:  # noqa: BLE001  没有主题就当作"无 CJK 主题"
        return ""


def _theme_has_cjk(xml: str) -> bool:
    """主题里是否已经声明了**东亚**字体。

    口径说明：python-pptx 默认模板写的是 `<a:font script="Hans" typeface="宋体"/>`，
    字面上"非空"，但那正是要消灭的宋体默认值 —— 宋体/SimSun 一律算**没有**。
    """
    m = re.search(r'<a:ea typeface="([^"]*)"', xml)
    if m and m.group(1).strip():
        return True
    m = re.search(r'<a:font script="Hans" typeface="([^"]*)"', xml)
    return bool(m and m.group(1).strip() and m.group(1) not in ("宋体", "SimSun"))


def read_template(path: str) -> TemplateInfo:
    """枚举模板的 slide_layouts 与占位符；不可用抛 PptxError("TEMPLATE_INVALID")。"""
    if not os.path.isfile(path) or not path.lower().endswith(".pptx"):
        raise PptxError(f"模板不可用：{os.path.basename(str(path))}", "TEMPLATE_INVALID",
                        "模板需为标准 .pptx，且至少含一个版式")
    try:
        prs = Presentation(path)
    except Exception as exc:  # noqa: BLE001
        raise PptxError(f"模板不可用：{os.path.basename(path)}（{exc}）",
                        "TEMPLATE_INVALID", "模板需为标准 .pptx，且至少含一个版式") from exc

    layouts = list(prs.slide_layouts)
    if not layouts:
        raise PptxError(f"模板不可用：{os.path.basename(path)}（没有任何版式）",
                        "TEMPLATE_INVALID", "模板需为标准 .pptx，且至少含一个版式")
    info = TemplateInfo(path=os.path.abspath(path),
                        layout_names=[ly.name for ly in layouts],
                        has_theme_cjk=_theme_has_cjk(_theme_xml(prs)))
    for i, layout in enumerate(layouts):
        info.placeholders[i] = [
            (ph.placeholder_format.idx, str(ph.placeholder_format.type), ph.name)
            for ph in layout.placeholders
        ]
    return info


def build_builtin_master(out_path: str, title: str = "演示文稿") -> str:
    """生成内置 5 版式的**空白母版**（真占位符），供用户在 PowerPoint 里手改。"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = builder.SLIDE_W, builder.SLIDE_H

    cover = prs.slides.add_slide(prs.slide_layouts[_ROLE_LAYOUT["cover"]])
    if not _set_ph_text(cover, _TITLE_PH, [(title, 0)], 40.0, True):
        _add_textbox(cover, _SAFE_LEFT, _SAFE_TOP, _SAFE_W, Inches(1.6),
                     [(title, 0)], 40.0, True)
    _set_ph_text(cover, _BODY_PH, [("副标题（可删）", 0)], 18.0, False)

    for role, label, body_text in (
            ("content", "内容页：标题 + 要点", [("要点一", 0), ("要点二", 0), ("要点三", 0)]),
            ("data", "数据页：标题 + 图表/表格", [("（此处放图表或表格）", 0)])):
        slide = prs.slides.add_slide(prs.slide_layouts[_ROLE_LAYOUT[role]])
        if not _set_ph_text(slide, _TITLE_PH, [(label, 0)], 28.0, True):
            _add_textbox(slide, _SAFE_LEFT, _SAFE_TOP, _SAFE_W, Inches(1.4),
                         [(label, 0)], 28.0, True)
        _set_ph_text(slide, _BODY_PH, body_text, 18.0, False)

    toc = prs.slides.add_slide(prs.slide_layouts[_ROLE_LAYOUT["content"]])
    if not _set_ph_text(toc, _TITLE_PH, [("目录", 0)], 28.0, True):
        _add_textbox(toc, _SAFE_LEFT, _SAFE_TOP, _SAFE_W, Inches(1.4),
                     [("目录", 0)], 28.0, True)
    _set_ph_text(toc, _BODY_PH, [("一、背景与现状", 0), ("二、核心问题", 0),
                                 ("三、结论与建议", 0)], 18.0, False)

    end = prs.slides.add_slide(prs.slide_layouts[_ROLE_LAYOUT["end"]])
    if not _set_ph_text(end, _TITLE_PH, [("谢谢观看", 0)], 40.0, True):
        _add_textbox(end, _SAFE_LEFT, _SAFE_TOP, _SAFE_W, Inches(1.6),
                     [("谢谢观看", 0)], 40.0, True)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    _patch_theme_zip(out_path)
    return out_path


# ---------------------------------------------------------------- 文本填充

def _set_run(run, size_pt: float, bold: bool, builtin: bool):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if builtin:
        run.font.name = FONT
        builder._set_ea(run, FONT)   # a:ea 必须排在 a:latin 之后


def _clean(text: str) -> str:
    """段落文本进一个 run —— 段内换行会破坏 run 语义，统一折成空格。"""
    return (text or "").replace("\n", " ").strip()


def _fill_tf(tf, paragraphs, size_pt: float, bold: bool, builtin: bool,
             width_pt: float | None, height_pt: float | None, report) -> None:
    """用 tf.clear() 留一个空段落再逐段写 —— 不要 tf.text = "…"（造不出字体控制）。"""
    tf.clear()
    tf.word_wrap = True
    first = True
    for text, level in paragraphs:
        text = _clean(text)
        if not text:
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        size = size_pt
        if width_pt and height_pt and width_pt > 0:
            size, text, truncated = fit_text(text, size_pt, width_pt, height_pt)
            if truncated and report is not None:
                report.setdefault("truncated", []).append(text)
        run = p.add_run()
        run.text = text
        _set_run(run, size, bold, builtin)


def _set_ph_text(slide, types, paragraphs, size_pt, bold, builtin=True,
                 report=None) -> bool:
    """按类型找占位符填文本；找不到记 degraded_placeholder 并返回 False。"""
    for ph in slide.placeholders:
        if ph.placeholder_format.type in types:
            tf = ph.text_frame
            w_pt = (ph.width - tf.margin_left - tf.margin_right) / EMU_PER_PT
            h_pt = (ph.height - tf.margin_top - tf.margin_bottom) / EMU_PER_PT
            _fill_tf(tf, paragraphs, size_pt, bold, builtin, w_pt, h_pt, report)
            return True
    if report is not None:
        report.setdefault("degraded_placeholder", []).append(
            f"{slide.slide_layout.name}:{types[0]}")
    return False


def _add_textbox(slide, left, top, width, height, paragraphs, size_pt, bold,
                 builtin=True, report=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    w_pt = (width - Emu(91440) * 2) / EMU_PER_PT
    h_pt = (height - Emu(45720) * 2) / EMU_PER_PT
    _fill_tf(box.text_frame, paragraphs, size_pt, bold, builtin, w_pt, h_pt, report)
    return box


# ---------------------------------------------------------------- 超长文本

def fit_text(text: str, size_pt: float, box_width_pt: float, box_height_pt: float,
             min_size_pt: float = 12.0) -> tuple:
    """先缩字号、再截断。返回 (最终字号, 最终文本, 是否被截断)。

    为什么缩字号优先：截断会**静默丢内容** —— CLI 的调用方是 agent，
    `{ok:true}` 里少一句话它无从察觉。缩字号保住信息，只有到 min_size_pt
    仍放不下（极端输入）才截断，且 truncated 必须冒泡给调用方。
    """
    def fits(t: str, s: float) -> bool:
        lines = qa.measure_text_lines(t, s, box_width_pt)
        return lines * s * qa.LINE_HEIGHT_FACTOR <= box_height_pt

    # 缩字号的**下限**：min_size_pt 与调用方给的字号取小 —— 调用方只想要 8pt 时
    # 不该被我们"放大"到 12pt（那会破版）。
    floor = min(float(size_pt), float(min_size_pt))
    size = float(size_pt)
    while size >= floor:
        if fits(text, size):
            return size, text, False
        size -= 1.0
    size = floor

    lo, hi = 1, len(text)
    best = "…"
    while lo <= hi:
        mid = (lo + hi) // 2
        cand = text[:mid] + "…"
        if fits(cand, size):
            best = cand
            lo = mid + 1
        else:
            hi = mid - 1
    return size, best, True


# ---------------------------------------------------------------- 主题后处理

def _replace_file(src: str, dst: str, attempts: int = 6) -> None:
    """把 src 覆盖到 dst，带退避重试与原地重写退路。

    Windows 上 `os.replace` 会偶发 `PermissionError: [WinError 5]`：刚写完的文件
    可能被杀软/索引器短暂持有。实测同一段代码同一目录，失败与否**每次不同**
    （`tools/probes/probe_zip_replace.py` 连跑两次，失败的用例换了位置）→
    是外部句柄的偶发占用，不是逻辑错。故：短暂退避重试；仍不行则退回
    "读全量 → 删原文件 → 原地重写"（不依赖 rename 的原子性，但保内容）。
    """
    last = None
    for i in range(attempts):
        try:
            os.replace(src, dst)
            return
        except PermissionError as exc:
            last = exc
            time.sleep(0.05 * (i + 1))
    try:
        with open(src, "rb") as f:
            data = f.read()
        if os.path.isfile(dst):
            os.remove(dst)
        with open(dst, "wb") as f:
            f.write(data)
        os.remove(src)
    except OSError:
        raise last


def _patch_theme_zip(path: str, font: str = FONT) -> bool:
    """把 theme*.xml 的东亚字体改成 font。返回是否真的改了。

    实现约束（契约 §7.4）：读原 zip → 新建临时 zip **全量复制**（只替换 theme）
    → 覆盖回去。绝不重排/丢条目——PowerPoint 对条目顺序不敏感，
    但对缺失条目零容忍。
    """
    tmp = path + ".tmp"
    changed = False
    try:
        with zipfile.ZipFile(path) as src:
            with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
                for item in src.infolist():
                    data = src.read(item.filename)
                    if re.fullmatch(r"ppt/theme/theme\d+\.xml", item.filename):
                        text = data.decode("utf-8")
                        new = text.replace('<a:ea typeface=""/>',
                                           f'<a:ea typeface="{font}"/>')
                        new = re.sub(r'(<a:font script="Hans" typeface=")[^"]*(")',
                                     rf"\g<1>{font}\g<2>", new)
                        if new != text:
                            data = new.encode("utf-8")
                            changed = True
                    dst.writestr(item, data)
        _replace_file(tmp, path)
    except Exception:
        if os.path.isfile(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
        raise
    return changed


# ---------------------------------------------------------------- 导出一本

def _delete_all_slides(prs) -> None:
    """模板自带的示例页要清掉——导出稿只应有我们生成的页。"""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)


def _layout_for(prs, role: str):
    layouts = list(prs.slide_layouts)
    if not layouts:
        raise PptxError("母版里没有任何版式", "TEMPLATE_INVALID", "换一份标准 .pptx 模板")
    return layouts[min(_ROLE_LAYOUT.get(role, 1), len(layouts) - 1)]


def _page_payload(deck, page, report):
    """把一页的讲解单元拆成 (标题, 正文段落, 表格文本)。"""
    units = hl_layout.build_units(hl_layout.page_shapes(page, deck))
    title, body, tables = None, [], []
    for u in units:
        if u.kind in ("picture", "chart"):
            # IR 只带几何、不带图片路径/图表数据 → 无法往返（契约缺陷节）
            report.setdefault("degraded", []).append(
                {"page": page["index"], "shape": u.shape_name, "kind": u.kind,
                 "reason": "ir_lacks_media"})
            continue
        if not u.text.strip():
            continue
        if u.kind in ("table", "cell"):
            tables.append(u.text)
        elif title is None and u.kind == "title":
            title = u.text
        else:
            body.append((u.text, 0))
    if title is None and body:
        title, body = body[0][0], body[1:]
    return title, body, tables


def build_deck_pptx(deck: DeckIR, out_path: str, template: str | None = None,
                    no_com: bool = False, report: dict | None = None) -> str:
    """deck.json → 可编辑 pptx。template=None 用内置母版；否则读自定义模板。

    `no_com=True`：不支持的形状跳过而非报错（本模块本就不使用 COM，该参数只为与
    CLI 契约对齐，语义是"允许降级而非抛错"）。
    """
    report = report if report is not None else {}
    info = read_template(template) if template else None
    try:
        prs = Presentation(template) if template else Presentation()
    except Exception as exc:  # noqa: BLE001
        raise PptxError(f"模板不可用：{template}（{exc}）", "TEMPLATE_INVALID",
                        "模板需为标准 .pptx，且至少含一个版式") from exc

    if info is None:
        # 内置母版：画布固定 16:9。**自定义模板不动画布**——那是品牌模板的意图
        prs.slide_width, prs.slide_height = builder.SLIDE_W, builder.SLIDE_H
    _delete_all_slides(prs)
    # 内置母版、或模板自带主题没有 CJK 字体时才覆盖字体；否则尊重品牌模板
    builtin = info is None or not info.has_theme_cjk
    report["font_overridden"] = bool(builtin)

    for i, page in enumerate(deck.pages):
        title, body, tables = _page_payload(deck, page, report)
        if i == 0:
            role = "cover"
        elif tables:
            role = "data"
        elif i == len(deck.pages) - 1 and not body:
            role = "end"
        else:
            role = "content"
        slide = prs.slides.add_slide(_layout_for(prs, role))

        if title:
            if not _set_ph_text(slide, _TITLE_PH, [(title, 0)],
                                32.0 if role == "cover" else 26.0, True,
                                builtin, report):
                _add_textbox(slide, _SAFE_LEFT, _SAFE_TOP, _SAFE_W, Inches(1.4),
                             [(title, 0)], 32.0, True, builtin, report)
        if body:
            if not _set_ph_text(slide, _BODY_PH, body, 18.0, False, builtin, report):
                _add_textbox(slide, _SAFE_LEFT, Inches(2.0), _SAFE_W, Inches(4.6),
                             body, 18.0, False, builtin, report)
        if tables:
            _add_table(slide, tables, builtin)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    if builtin:
        _patch_theme_zip(out_path)
    return out_path


def _add_table(slide, texts, builtin):
    """表格用 add_table 落在正文区（PowerPoint 里可编辑）。"""
    rows = len(texts)
    width, height = int(_SAFE_W), int(Inches(4.4))
    shape = slide.shapes.add_table(rows, 1, int(_SAFE_LEFT), int(Inches(2.0)),
                                   width, height)
    table = shape.table
    for i, text in enumerate(texts):
        cell = table.cell(i, 0)
        cell.text_frame.clear()
        run = cell.text_frame.paragraphs[0].add_run()
        run.text = _clean(text)
        _set_run(run, 16.0, False, builtin)
