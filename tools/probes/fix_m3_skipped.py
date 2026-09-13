"""M3 对照探针：`build_units` 丢弃形状时有没有出口。

背景（AUDIT_REPORT §2 M3）：契约 §4.3 明写 `inner_w_pt <= 0` → "跳过该形状（**记
warning**，不产出 Unit）"。实现只做到"不产出 Unit"：`return` 一句就走了，调用方
**没有任何通道**知道"这里为什么少了一块高亮"。`read_pages` 有 `skipped` 出口，
布局期丢弃却一条都进不去 CLI 的 JSON。

本探针纯内存/临时文件，不启动任何外部程序（无需 COM 闸门）。

判定：
  修复前 → `build_units` 没有收集器参数（传就 TypeError）；CLI 的 warnings 里看不到
  修复后 → 收集器拿到 {page_index, shape_id, shape_name, reason}；CLI warnings 里可见

跑法：.venv/Scripts/python.exe tools/probes/fix_m3_skipped.py
"""

import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, ROOT)

# 一个内宽为负的文本框：宽 100000 EMU，左右内边距各 91440 EMU（合计 182880）
NARROW_SHAPE_W = 100000
MARGIN_LR = 91440


def build_case_deck(path: str) -> str:
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ok = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    ok.name = "NormalBox"
    ok.text_frame.text = "正常形状"
    ok.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    bad = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(NARROW_SHAPE_W), Emu(400000))
    bad.name = "TooNarrowBox"
    bad.text_frame.text = "边距吃掉整宽"
    prs.save(path)
    return path


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="m3probe_")
    try:
        src = build_case_deck(os.path.join(tmp, "case.pptx"))
        import hl_layout
        import pptx_io

        deck, _ = pptx_io.read_pages(src)
        page = hl_layout.page_shapes(deck.pages[0], deck)

        inner_w_pt = (NARROW_SHAPE_W - 2 * MARGIN_LR) / 12700
        print(f"样本：{src}")
        print(f"  窄框：宽 {NARROW_SHAPE_W} EMU，左右边距各 {MARGIN_LR} EMU "
              f"→ inner_w_pt = {inner_w_pt:.2f}（<= 0）")

        print("\n== 1) 不带收集器 ==")
        units = hl_layout.build_units(page)
        texts = [u.text for u in units]
        print(f"  units = {len(units)}，文本 = {texts}")
        print("  → 窄框不见了，但**没有任何理由**告诉调用方（M3 的现象）")

        print("\n== 2) 带收集器 skipped=[] ==")
        skipped: list = []
        try:
            units2 = hl_layout.build_units(page, skipped=skipped)
        except TypeError as exc:
            print(f"  **不支持收集器** → TypeError: {exc}")
            print("  判定：**REPRODUCED** —— 丢弃是静默的，没有出口")
            return 0
        print(f"  units = {len(units2)}，skipped = {skipped}")

        print("\n== 3) 走 CLI：warnings 里能不能看到 ==")
        import io
        import json
        from contextlib import redirect_stdout

        import cli
        import outline
        # redesign 模式会调 LLM；本探针只关心 warnings 的出口，把流水线打桩掉
        outline.generate_outline_from_text = lambda text, density="balanced": []
        cli._design_pipeline = lambda out, topic, slides, image_map=None: ""

        buf = io.StringIO()
        out_dir = os.path.join(tmp, "out")
        with redirect_stdout(buf):
            rc = cli.main(["import", src, "--no-com", "--mode", "redesign",
                           "--out", out_dir])
        payload = json.loads(buf.getvalue())
        data = payload.get("data", {})
        warnings = data.get("warnings", [])
        print(f"  CLI 退出码 = {rc}")
        print(f"  skipped = {data.get('skipped')}，warnings = {warnings}")

        ok = bool(skipped) and any("no_inner_width" in w for w in warnings)
        print("\n" + "-" * 68)
        print(f"判定：{'**FIXED** —— 丢弃有出口，CLI warnings 里可见' if ok else '**未修复**'}")
        print("-" * 68)
        return 0 if ok else 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
