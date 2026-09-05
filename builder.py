"""第三步：python-pptx 按主题色板 + 多版式组装 PPT（色板来自 ppt-maker skill）。"""

import os

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "微软雅黑"

# 主题色板来自 ppt-maker skill（~/.agents/skills/ppt-maker）
THEMES = {
    "blue":  {"bg": RGBColor(0x0F, 0x2A, 0x4A), "accent": RGBColor(0x3B, 0x82, 0xF6),
              "fg": RGBColor(0xFF, 0xFF, 0xFF), "muted": RGBColor(0xB6, 0xC7, 0xDC)},
    "dark":  {"bg": RGBColor(0x11, 0x18, 0x27), "accent": RGBColor(0xF5, 0x9E, 0x0B),
              "fg": RGBColor(0xF9, 0xFA, 0xFB), "muted": RGBColor(0x9C, 0xA3, 0xAF)},
    "green": {"bg": RGBColor(0x06, 0x2E, 0x21), "accent": RGBColor(0x22, 0xC5, 0x5E),
              "fg": RGBColor(0xEC, 0xFD, 0xF5), "muted": RGBColor(0xA7, 0xC4, 0xB6)},
}

# 图表类型映射（chart.type → XL_CHART_TYPE）
_CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,        # 横向条形
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,  # 纵向柱状
    "pie": XL_CHART_TYPE.PIE,                  # 饼图
    "line": XL_CHART_TYPE.LINE_MARKERS,        # 折线
}


def _fit_title_size(title: str) -> int:
    """标题按长度自适应字号，避免溢出。"""
    n = len(title)
    if n <= 14:
        return 28
    if n <= 20:
        return 24
    if n <= 28:
        return 20
    return 18


def _fit_body_size(points) -> int:
    """要点按总长度自适应字号。"""
    total = sum(len(p) for p in points)
    if total <= 60:
        return 17
    if total <= 110:
        return 15
    return 13


def _set_ea(run, font=FONT):
    """给 run 补写 a:ea（东亚字体）。

    python-pptx 的 run.font.name 只写 a:latin，中文属于东亚文字范围，机器上按
    a:ea 匹配字体；缺 a:ea 时中文会回退用户机器的默认字体（如宋体），版式失真。
    a:ea 在 rPr 子元素里必须排在 a:latin 之后——font.name 已先写好 a:latin，
    且 run 的 rPr 中不存在 a:cs/sym/hlink 等后续元素，这里直接 append 即满足顺序。
    """
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)


def _set_text(frame, text, size, color, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    frame.auto_size = MSO_AUTO_SIZE.NONE  # 显式禁用自动调整，防止 PowerPoint 打开时文本框自动长高
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    _set_ea(run, font)
    return p


def _add_bg(slide, color):
    rect = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()
    return rect


def _add_title(slide, title, t):
    """标题 + 左侧强调色竖条，字号按长度自适应。"""
    bar = slide.shapes.add_shape(1, Inches(0.6), Inches(0.78), Inches(0.09), Inches(0.52))
    bar.fill.solid()
    bar.fill.fore_color.rgb = t["accent"]
    bar.line.fill.background()
    _set_text(slide.shapes.add_textbox(Inches(0.85), Inches(0.7), Inches(6.5), Inches(0.9)).text_frame,
              title, _fit_title_size(title), t["fg"], bold=True)


def _add_cover(slide, title, subtitle, t):
    _set_text(slide.shapes.add_textbox(Inches(1), Inches(2.6), SLIDE_W - Inches(2), Inches(2)).text_frame,
              title, 44, t["fg"], bold=True, align=PP_ALIGN.CENTER)
    if subtitle:
        _set_text(slide.shapes.add_textbox(Inches(1), Inches(3.9), SLIDE_W - Inches(2), Inches(0.8)).text_frame,
                  subtitle, 18, t["muted"], align=PP_ALIGN.CENTER)


def _add_toc(slide, title, items, t):
    _add_title(slide, title, t)
    box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.0), Inches(4.8))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(20)
        run = p.add_run()
        run.text = f"{i + 1:>2}    {it}"
        run.font.name = FONT
        run.font.size = Pt(22)
        run.font.color.rgb = t["fg"] if i == 0 else t["muted"]
        run.font.bold = (i == 0)
        _set_ea(run)


def _add_section(slide, title, points, t):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.7), SLIDE_W - Inches(1.6), Inches(1.5))
    _set_text(box.text_frame, title, 42, t["fg"], bold=True, align=PP_ALIGN.CENTER)
    if points:
        _set_text(slide.shapes.add_textbox(Inches(0.8), Inches(4.2), SLIDE_W - Inches(1.6), Inches(0.8)).text_frame,
                  points[0], 18, t["muted"], align=PP_ALIGN.CENTER)
    line = slide.shapes.add_shape(1, Inches(3.5), Inches(4.0), Inches(6.3), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = t["accent"]
    line.line.fill.background()


def _split_point(pt: str):
    """拆两级：'标题：描述' → (标题, 描述)；无冒号 → (原文, '')。"""
    for sep in ("：", ":"):
        if sep in pt:
            title, desc = pt.split(sep, 1)
            return title.strip(), desc.strip()
    return pt.strip(), ""


def _resolve_layout(layout, points, has_image):
    """layout 有值用，否则按内容特征兜底决策。"""
    if layout:
        return layout
    n = len(points)
    if has_image:
        return "image-top" if n >= 4 else "image-right"
    if n <= 2:
        return "center"
    if n == 3:
        return "cards"
    return "columns"


def _add_image(slide, image_path, x, y, w, h) -> bool:
    """等比缩放图片放入 (x,y,w,h) 区域，成功返回 True。"""
    if not (image_path and os.path.exists(image_path)):
        return False
    with Image.open(image_path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    dw, dh = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(image_path, x + int((w - dw) / 2), y + int((h - dh) / 2),
                             width=Emu(dw), height=Emu(dh))
    return True


def _add_point_list(slide, points, left, top, width, height, t, start=0, title_size=16, desc_size=13):
    """渲染两级要点（编号 + 标题 accent 加粗 + 描述 muted 缩进）。"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, pt in enumerate(points):
        title, desc = _split_point(pt)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(20)
        run = p.add_run()
        run.text = f"{start + i + 1:02d}  {title}"
        run.font.name = FONT
        run.font.size = Pt(title_size)
        run.font.bold = True
        run.font.color.rgb = t["accent"]
        _set_ea(run)
        if desc:
            p2 = tf.add_paragraph()
            p2.space_before = Pt(2)
            run2 = p2.add_run()
            run2.text = "      " + desc
            run2.font.name = FONT
            run2.font.size = Pt(desc_size)
            run2.font.color.rgb = t["muted"]
            _set_ea(run2)
    return box


def _add_content_image_right(slide, title, points, image_path, t):
    _add_title(slide, title, t)
    _add_point_list(slide, points, Inches(0.6), Inches(1.8), Inches(6.2), Inches(4.9), t)
    _add_image(slide, image_path, Inches(7.2), Inches(0.9), Inches(5.5), Inches(5.7))


def _add_content_image_left(slide, title, points, image_path, t):
    _add_image(slide, image_path, Inches(0.6), Inches(0.9), Inches(5.5), Inches(5.7))
    bar = slide.shapes.add_shape(1, Inches(7.2), Inches(0.8), Inches(0.09), Inches(0.52))
    bar.fill.solid()
    bar.fill.fore_color.rgb = t["accent"]
    bar.line.fill.background()
    _set_text(slide.shapes.add_textbox(Inches(7.45), Inches(0.7), Inches(5.3), Inches(0.9)).text_frame,
              title, _fit_title_size(title), t["fg"], bold=True)
    _add_point_list(slide, points, Inches(7.2), Inches(1.8), Inches(5.5), Inches(4.9), t)


def _add_content_image_top(slide, title, points, image_path, t):
    _add_title(slide, title, t)
    _add_image(slide, image_path, Inches(0.6), Inches(1.7), Inches(12.0), Inches(2.8))
    _add_point_list(slide, points, Inches(0.6), Inches(4.6), Inches(12.0), Inches(2.5), t)


def _add_content_image_full(slide, title, points, image_path, t):
    _add_image(slide, image_path, 0, 0, SLIDE_W, SLIDE_H)
    card = slide.shapes.add_shape(1, Inches(0.6), Inches(1.2), Inches(6.2), Inches(5.0))
    card.fill.solid()
    card.fill.fore_color.rgb = t["bg"]
    card.line.color.rgb = t["accent"]
    card.line.width = Pt(1.5)
    _set_text(slide.shapes.add_textbox(Inches(0.95), Inches(1.4), Inches(5.5), Inches(0.9)).text_frame,
              title, 26, t["fg"], bold=True)
    _add_point_list(slide, points, Inches(0.95), Inches(2.3), Inches(5.5), Inches(3.7), t)


def _add_content_cards(slide, title, points, image_path, t):
    _add_title(slide, title, t)
    n = min(len(points), 4)
    gap = Inches(0.4)
    total = Inches(11.6)
    card_w = int((total - gap * (n - 1)) / n)
    for i in range(n):
        title_text, desc = _split_point(points[i])
        x = Inches(0.7) + (card_w + gap) * i
        card = slide.shapes.add_shape(5, x, Inches(2.1), card_w, Inches(4.2))
        card.fill.solid()
        card.fill.fore_color.rgb = t["bg"]
        card.line.color.rgb = t["accent"]
        card.line.width = Pt(1.25)
        top = slide.shapes.add_shape(1, x, Inches(2.1), card_w, Inches(0.07))
        top.fill.solid()
        top.fill.fore_color.rgb = t["accent"]
        top.line.fill.background()
        _set_text(slide.shapes.add_textbox(x + Inches(0.3), Inches(2.4), card_w - Inches(0.6), Inches(0.6)).text_frame,
                  f"{i + 1:02d}", 26, t["accent"], bold=True)
        box = slide.shapes.add_textbox(x + Inches(0.3), Inches(3.1), card_w - Inches(0.6), Inches(3.0))
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        run = tf.paragraphs[0].add_run()
        run.text = title_text
        run.font.name = FONT
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = t["fg"]
        _set_ea(run)
        if desc:
            p2 = tf.add_paragraph()
            p2.space_before = Pt(6)
            run2 = p2.add_run()
            run2.text = desc
            run2.font.name = FONT
            run2.font.size = Pt(12)
            run2.font.color.rgb = t["muted"]
            _set_ea(run2)


def _add_content_columns(slide, title, points, image_path, t):
    _add_title(slide, title, t)
    mid = (len(points) + 1) // 2
    _add_point_list(slide, points[:mid], Inches(0.6), Inches(1.8), Inches(5.8), Inches(4.9), t, start=0)
    _add_point_list(slide, points[mid:], Inches(6.8), Inches(1.8), Inches(5.8), Inches(4.9), t, start=mid)


def _add_content_center(slide, title, points, image_path, t):
    _set_text(slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.3), Inches(1.2)).text_frame,
              title, 40, t["fg"], bold=True, align=PP_ALIGN.CENTER)
    box = slide.shapes.add_textbox(Inches(1.5), Inches(3.8), Inches(10.3), Inches(2.4))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, pt in enumerate(points):
        title_text, desc = _split_point(pt)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(18)
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title_text + (f"：{desc}" if desc else "")
        run.font.name = FONT
        run.font.size = Pt(22)
        run.font.color.rgb = t["muted"]
        _set_ea(run)


_CONTENT_LAYOUTS = {
    "image-right": _add_content_image_right,
    "image-left": _add_content_image_left,
    "image-top": _add_content_image_top,
    "image-full": _add_content_image_full,
    "cards": _add_content_cards,
    "columns": _add_content_columns,
    "center": _add_content_center,
}


def _add_content(slide, title, points, image_path, t, layout=None):
    has_image = bool(image_path and os.path.exists(image_path))
    resolved = _resolve_layout(layout, points, has_image)
    _CONTENT_LAYOUTS.get(resolved, _add_content_image_right)(slide, title, points, image_path, t)


def _pick_chart_type(chart) -> str:
    """显式 chart.type 优先，否则按数据特征自动选型。"""
    ct = chart.get("type")
    if ct in _CHART_TYPES:
        return ct
    labels = chart.get("labels", [])
    values = chart.get("values", [])
    # labels 像时间序列（含年份数字）→ 折线
    if labels and all(isinstance(l, str) and any(c.isdigit() for c in l) for l in labels):
        return "line"
    # 值接近占比（总和≈100 且项少）→ 饼图
    if values and len(values) <= 6 and abs(sum(values) - 100) < 30:
        return "pie"
    return "column"


def _add_data(slide, title, chart, t):
    _add_title(slide, title, t)
    ctype = _pick_chart_type(chart)
    labels = chart.get("labels", [])
    values = chart.get("values", [])
    chart_data = CategoryChartData()
    chart_data.categories = labels
    chart_data.add_series("数值", values)
    gframe = slide.shapes.add_chart(_CHART_TYPES[ctype], Inches(0.6), Inches(1.9),
                                    Inches(12.0), Inches(5.1), chart_data)
    chart_obj = gframe.chart
    chart_obj.has_legend = (ctype == "pie")
    try:
        if ctype != "pie":
            chart_obj.value_axis.has_major_gridlines = False
            chart_obj.category_axis.tick_labels.font.color.rgb = t["fg"]
            chart_obj.value_axis.tick_labels.font.color.rgb = t["muted"]
            chart_obj.category_axis.tick_labels.font.size = Pt(11)
            chart_obj.value_axis.tick_labels.font.size = Pt(10)
    except Exception:
        pass
    if chart_obj.plots and chart_obj.plots[0].series:
        series = chart_obj.plots[0].series[0]
        try:
            if ctype == "pie":
                series.data_labels.font.color.rgb = t["fg"]
                series.data_labels.font.size = Pt(11)
                series.data_labels.number_format = "0"
                series.data_labels.number_format_is_linked = False
                chart_obj.plots[0].has_data_labels = True
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = t["accent"]
                series.data_labels.font.color.rgb = t["fg"]
                series.data_labels.font.size = Pt(10)
                series.data_labels.number_format = "0"
                series.data_labels.number_format_is_linked = False
                chart_obj.plots[0].has_data_labels = True
        except Exception:
            pass


def _add_timeline(slide, title, points, t):
    _add_title(slide, title, t)
    n = len(points)
    if n == 0:
        return
    # 横向时间轴
    line = slide.shapes.add_shape(1, Inches(0.8), Inches(3.3), Inches(11.7), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = t["accent"]
    line.line.fill.background()
    step = 11.7 / n
    for i, ev in enumerate(points):
        cx = Inches(0.8 + step * (i + 0.5))
        dot = slide.shapes.add_shape(9, int(cx) - int(Inches(0.1)), Inches(3.22), Inches(0.2), Inches(0.2))
        dot.fill.solid()
        dot.fill.fore_color.rgb = t["accent"]
        dot.line.fill.background()
        # 事件文字（节点下方，小字）
        _set_text(slide.shapes.add_textbox(Inches(0.4 + step * i), Inches(3.7), Inches(step), Inches(1.6)).text_frame,
                  ev, 13, t["fg"], align=PP_ALIGN.CENTER)


def _add_compare(slide, title, points, t):
    _add_title(slide, title, t)
    mid = (len(points) + 1) // 2
    left = points[:mid]
    right = points[mid:]
    for col, items in ((0, left), (1, right)):
        x = Inches(0.8) + Inches(6.2) * col
        # 栏标题条
        bar = slide.shapes.add_shape(1, x, Inches(2.0), Inches(5.5), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = t["accent"]
        bar.line.fill.background()
        box = slide.shapes.add_textbox(x, Inches(2.2), Inches(5.5), Inches(4.0))
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(14)
            run = p.add_run()
            run.text = f"▸ {it}"
            run.font.name = FONT
            run.font.size = Pt(16)
            run.font.color.rgb = t["muted"]
            _set_ea(run)


def _add_end(slide, title, t):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), SLIDE_W - Inches(1.6), Inches(1.6))
    _set_text(box.text_frame, title, 36, t["fg"], bold=True, align=PP_ALIGN.CENTER)


def build_ppt(slides: list[dict], image_paths: list[str | None], out_path: str,
              theme: str = "blue", subtitle: str = "") -> str:
    t = THEMES.get(theme, THEMES["blue"])
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for i, s in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        stype = s.get("type", "")
        title = s.get("title", "")
        points = s.get("points", [])
        _add_bg(slide, t["bg"])

        if stype == "cover" or (i == 0 and not points):
            _add_cover(slide, title, subtitle or s.get("image_prompt", ""), t)
        elif stype == "toc":
            _add_toc(slide, title, points, t)
        elif stype == "section":
            _add_section(slide, title, points, t)
        elif stype == "timeline":
            _add_timeline(slide, title, points, t)
        elif stype == "compare":
            _add_compare(slide, title, points, t)
        elif stype == "data" and s.get("chart") and s["chart"].get("values"):
            _add_data(slide, title, s["chart"], t)
        elif stype == "end":
            _add_end(slide, title, t)
        else:  # content / 未知 type
            _add_content(slide, title, points,
                         image_paths[i] if i < len(image_paths) else None, t, s.get("layout"))

    prs.save(out_path)
    return out_path
