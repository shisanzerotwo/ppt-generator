"""M6：`pptx_out.py` 单测 —— 模板枚举、母版生成、可编辑导出、中文字体、超长文本。

字体这块分两层锁：
- **静态**：theme1.xml 的 `a:ea`/`<a:font script="Hans">` 与每个 run 的 `a:ea`
  必须是微软雅黑，且 `a:ea` 排在 `a:latin` **之后**（OOXML 元素序要求）；
- **可编辑**：导出稿改字能存回来、页数不变（占位符是真占位符，不是画上去的图片）。

"渲染成微软雅黑而非宋体"的端到端对照需要 COM，放在
`tools/probes/accept_m6.py`（含阳性对照），单测只锁机制层。
"""

import os
import re
import zipfile

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

import pptx_io
import pptx_out
import qa
from pptx_io import PptxError

FONT = "微软雅黑"


# ---------------------------------------------------------------- 造稿

def _deck(tmp_path, pages=3, titles=None):
    """用 python-pptx 造一份稿再读成 DeckIR（不碰 COM）。"""
    path = tmp_path / "src.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i in range(pages):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        box.text_frame.text = (titles[i] if titles else f"第{i + 1}页标题")
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
        body = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(3))
        body.text_frame.text = f"第{i + 1}页要点甲"
        p2 = body.text_frame.add_paragraph()
        p2.add_run().text = f"第{i + 1}页要点乙"
        p2.runs[0].font.size = Pt(18)
        body.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    prs.save(str(path))
    deck, _ = pptx_io.read_pages(str(path), export_width_px=1920)
    return deck


def _theme_facts(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        theme = z.read("ppt/theme/theme1.xml").decode("utf-8")
        ea = re.findall(r'<a:ea typeface="([^"]*)"', theme)
        hans = re.findall(r'<a:font script="Hans" typeface="([^"]*)"', theme)
        order_ok = total = 0
        for name in names:
            if not re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
                continue
            xml = z.read(name).decode("utf-8")
            for rPr in re.findall(r"<a:rPr[^>]*>.*?</a:rPr>", xml, re.S):
                if "<a:ea" not in rPr:
                    continue
                total += 1
                if "<a:latin" in rPr and rPr.index("<a:latin") < rPr.index("<a:ea"):
                    order_ok += 1
        return ea, hans, order_ok, total, len(names)


# ---------------------------------------------------------------- 模板

def test_read_template_lists_layouts_and_placeholders(tmp_path):
    path = tmp_path / "brand.pptx"
    Presentation().save(str(path))
    info = pptx_out.read_template(str(path))
    assert len(info.layout_names) >= 6
    assert 0 in info.placeholders and 1 in info.placeholders
    assert any("TITLE" in str(t) for _, t, _ in info.placeholders[0])


def test_read_template_default_theme_has_no_cjk(tmp_path):
    """python-pptx 默认模板 `a:ea=""` + Hans=宋体 → 必须判成"没有 CJK 主题"。"""
    path = tmp_path / "brand.pptx"
    Presentation().save(str(path))
    assert pptx_out.read_template(str(path)).has_theme_cjk is False


def test_read_template_accepts_cjk_theme(tmp_path):
    """把主题的 Hans 改成非宋体后，应判成"已有 CJK 主题"（不再覆盖品牌字体）。"""
    src = tmp_path / "a.pptx"
    Presentation().save(str(src))
    dst = tmp_path / "b.pptx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/theme/theme1.xml":
                data = data.decode("utf-8").replace(
                    'script="Hans" typeface="宋体"',
                    'script="Hans" typeface="思源黑体"').encode("utf-8")
            zout.writestr(item, data)
    assert pptx_out.read_template(str(dst)).has_theme_cjk is True


@pytest.mark.parametrize("bad", ["nope.pptx", "brand.txt"])
def test_read_template_invalid(tmp_path, bad):
    p = tmp_path / bad
    if bad.endswith(".txt"):
        p.write_text("x", encoding="utf-8")
    with pytest.raises(PptxError) as ei:
        pptx_out.read_template(str(p))
    assert ei.value.code == "TEMPLATE_INVALID"


def test_read_template_rejects_corrupt_pptx(tmp_path):
    p = tmp_path / "broken.pptx"
    p.write_bytes(b"not a zip")
    with pytest.raises(PptxError) as ei:
        pptx_out.read_template(str(p))
    assert ei.value.code == "TEMPLATE_INVALID"


# ---------------------------------------------------------------- 内置母版

def test_build_builtin_master_five_layouts(tmp_path):
    out = tmp_path / "master.pptx"
    pptx_out.build_builtin_master(str(out), title="我的母版")
    prs = Presentation(str(out))
    assert len(prs.slides) == 5
    names = [s.slide_layout.name for s in prs.slides]
    assert "Title Slide" in names and "Title and Content" in names
    ea, hans, _, _, _ = _theme_facts(str(out))
    assert ea and all(v == FONT for v in ea)
    assert hans and all(v == FONT for v in hans)


# ---------------------------------------------------------------- 导出

def test_build_deck_pptx_page_count_and_no_errors(tmp_path):
    deck = _deck(tmp_path, pages=3)
    out = tmp_path / "out.pptx"
    pptx_out.build_deck_pptx(deck, str(out))
    assert len(Presentation(str(out)).slides) == 3
    check = qa.check_pptx(str(out), 3)
    assert check["errors"] == []


def test_build_deck_pptx_theme_font_is_yahei(tmp_path):
    deck = _deck(tmp_path, pages=2)
    out = tmp_path / "out.pptx"
    pptx_out.build_deck_pptx(deck, str(out))
    ea, hans, _, _, _ = _theme_facts(str(out))
    assert ea == [FONT, FONT], f"theme a:ea 应是微软雅黑：{ea}"
    assert hans == [FONT, FONT], f"theme Hans 应是微软雅黑：{hans}"


def test_build_deck_pptx_runs_have_ea_after_latin(tmp_path):
    """每个 run 补 a:ea，且**排在 a:latin 之后**（元素序不合规 PowerPoint 会修包）。"""
    deck = _deck(tmp_path, pages=2)
    out = tmp_path / "out.pptx"
    pptx_out.build_deck_pptx(deck, str(out))
    _, _, order_ok, total, _ = _theme_facts(str(out))
    assert total > 0
    assert order_ok == total, f"{order_ok}/{total} 的 a:ea 位置不对"


def test_build_deck_pptx_theme_patch_preserves_all_entries(tmp_path):
    """zip 后处理只换 theme，条目必须一个不少。"""
    deck = _deck(tmp_path, pages=1)
    out = tmp_path / "out.pptx"
    pptx_out.build_deck_pptx(deck, str(out))
    with zipfile.ZipFile(str(out)) as z:
        names = set(z.namelist())
    assert "[Content_Types].xml" in names
    assert "ppt/presentation.xml" in names
    assert "docProps/app.xml" in names or "docProps/core.xml" in names
    assert not any(n.endswith(".tmp") for n in names)


def test_build_deck_pptx_is_editable(tmp_path):
    """真占位符：改字能存回来、页数不变。"""
    deck = _deck(tmp_path, pages=2, titles=["原标题", "第二页"])
    out = tmp_path / "out.pptx"
    pptx_out.build_deck_pptx(deck, str(out))
    prs = Presentation(str(out))
    shape = next(s for s in prs.slides[0].shapes
                 if s.has_text_frame and s.text_frame.text.strip())
    shape.text_frame.paragraphs[0].runs[0].text = "改过的"
    edited = tmp_path / "edited.pptx"
    prs.save(str(edited))
    back = Presentation(str(edited))
    assert len(back.slides) == 2
    assert "改过的" in back.slides[0].shapes[0].text_frame.text


def test_build_deck_pptx_leaves_no_temp_files(tmp_path):
    deck = _deck(tmp_path, pages=1)
    out = tmp_path / "out.pptx"
    pptx_out.build_deck_pptx(deck, str(out))
    leftovers = [f for f in os.listdir(tmp_path) if f.endswith(".tmp")]
    assert leftovers == []


def test_media_shapes_are_degraded_loudly(tmp_path):
    """picture/chart 往返不了（IR 不带媒体），必须进 report 而不是静默消失。"""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from PIL import Image
    img = tmp_path / "p.png"
    Image.new("RGB", (40, 30), (200, 100, 50)).save(str(img))

    src = tmp_path / "withmedia.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "带媒体的页"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    slide.shapes.add_picture(str(img), Inches(1), Inches(3), Inches(2), Inches(1.5))
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("s", (1.0, 2.0))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(3),
                           Inches(4), Inches(2.5), data)
    prs.save(str(src))

    deck, _ = pptx_io.read_pages(str(src))
    report = {}
    pptx_out.build_deck_pptx(deck, str(tmp_path / "o.pptx"), report=report)
    kinds = {d["kind"] for d in report.get("degraded", [])}
    assert kinds == {"picture", "chart"}, report


def test_custom_template_with_cjk_theme_keeps_its_font(tmp_path):
    """品牌模板已声明 CJK 字体时不覆盖（尊重模板意图）。"""
    tpl = tmp_path / "brand.pptx"
    Presentation().save(str(tpl))
    patched = tmp_path / "brand2.pptx"
    with zipfile.ZipFile(str(tpl)) as zin, zipfile.ZipFile(str(patched), "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/theme/theme1.xml":
                data = data.decode("utf-8").replace(
                    'script="Hans" typeface="宋体"',
                    'script="Hans" typeface="思源黑体"').encode("utf-8")
            zout.writestr(item, data)

    deck = _deck(tmp_path, pages=1)
    out = tmp_path / "o.pptx"
    report = {}
    pptx_out.build_deck_pptx(deck, str(out), template=str(patched), report=report)
    assert report["font_overridden"] is False
    assert "思源黑体" in zipfile.ZipFile(str(out)).read(
        "ppt/theme/theme1.xml").decode("utf-8")


def test_template_sample_slides_are_removed(tmp_path):
    """模板自带的示例页要清掉，导出稿只应有我们生成的页。"""
    tpl = tmp_path / "brand.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(str(tpl))

    deck = _deck(tmp_path, pages=2)
    out = tmp_path / "o.pptx"
    pptx_out.build_deck_pptx(deck, str(out), template=str(tpl))
    assert len(Presentation(str(out)).slides) == 2


def test_custom_template_keeps_its_canvas_size(tmp_path):
    """自定义模板的画布尺寸不能被我们改成 16:9。"""
    tpl = tmp_path / "four3.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    prs.save(str(tpl))
    deck = _deck(tmp_path, pages=1)
    out = tmp_path / "o.pptx"
    pptx_out.build_deck_pptx(deck, str(out), template=str(tpl))
    got = Presentation(str(out))
    assert got.slide_width == Inches(10)


# ---------------------------------------------------------------- fit_text

def test_fit_text_keeps_size_when_it_fits():
    size, text, truncated = pptx_out.fit_text("短文本", 18.0, 400.0, 200.0)
    assert (size, text, truncated) == (18.0, "短文本", False)


def test_fit_text_shrinks_before_truncating():
    """优先缩字号：内容一个字都不能少。"""
    long_text = "人工智能" * 5          # 20 字；36pt 放不进 200×60，20pt 放得下
    size, text, truncated = pptx_out.fit_text(long_text, 36.0, 200.0, 60.0)
    assert size < 36.0, "应当缩小字号"
    assert truncated is False
    assert text == long_text, "缩字号能放下就不该截断"


def test_fit_text_truncates_at_floor_with_ellipsis():
    size, text, truncated = pptx_out.fit_text("人工智能" * 500, 32.0, 120.0, 40.0,
                                              min_size_pt=12.0)
    assert truncated is True
    assert size == pytest.approx(12.0)
    assert text.endswith("…") and len(text) < 100


def test_fit_text_never_enlarges_below_min_size():
    """调用方只要 8pt 时不该被"放大"到 min_size_pt（那会破版）。"""
    size, _, _ = pptx_out.fit_text("字" * 200, 8.0, 50.0, 10.0, min_size_pt=12.0)
    assert size <= 8.0


# ---------------------------------------------------------------- 脆弱环节

def test_replace_file_falls_back_when_rename_is_blocked(tmp_path, monkeypatch):
    """Windows 上 os.replace 会被偶发占用挡掉（PermissionError）→ 必须能自愈。"""
    src = tmp_path / "a.tmp"
    dst = tmp_path / "a.bin"
    src.write_bytes(b"NEW")
    dst.write_bytes(b"OLD")

    def boom(a, b):
        raise PermissionError(5, "拒绝访问")

    monkeypatch.setattr(pptx_out.os, "replace", boom)
    pptx_out._replace_file(str(src), str(dst), attempts=1)
    assert dst.read_bytes() == b"NEW"
    assert not src.exists()


def test_replace_file_retries_then_succeeds(tmp_path, monkeypatch):
    calls = {"n": 0}
    real = os.replace

    def flaky(a, b):
        calls["n"] += 1
        if calls["n"] < 3:
            raise PermissionError(5, "拒绝访问")
        return real(a, b)

    src = tmp_path / "a.tmp"
    dst = tmp_path / "a.bin"
    src.write_bytes(b"NEW")
    monkeypatch.setattr(pptx_out.os, "replace", flaky)
    pptx_out._replace_file(str(src), str(dst))
    assert calls["n"] == 3
    assert dst.read_bytes() == b"NEW"
