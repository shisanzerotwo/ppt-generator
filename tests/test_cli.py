"""M4：`pptgen` CLI 单测。

重点锁三件事：
1. **进程契约**：stdout 恒为一行合法 JSON（成功与失败都成立），人类日志只进 stderr；
2. **退出码**：0/2/3/4/5 与契约 §9 的 code 对应；
3. **明确的输入校验**：非法参数组合与病态输入一律走 PptxError，不给裸堆栈。

需要 PowerPoint / 浏览器的用例照抄 `test_shot_settle.py` 的 skipif 模式。
"""

import json
import os
import subprocess
import sys

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

import cli
import hl_anim
import pptx_io
import video

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


# ---------------------------------------------------------------- 工具

def _run(argv, capsys):
    """跑一次 CLI，返回 (退出码, stdout 解析出的 JSON, stderr 文本)。"""
    code = cli.main(argv)
    captured = capsys.readouterr()
    payload = json.loads(captured.out)  # 只允许一行、且必须可解析
    return code, payload, captured.err


def _assert_single_line_json(out: str):
    assert out.endswith("\n")
    assert out.count("\n") == 1, f"stdout 必须只有一行 JSON，实际 {out.count(chr(10))} 行"


def _src_pptx(tmp_path, pages=2):
    path = tmp_path / "src.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for i in range(pages):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.name = f"TitleBox{i}"
        box.text_frame.text = f"第{i + 1}页标题"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    prs.save(str(path))
    return str(path)


def _deck_dir(tmp_path, colors=((200, 30, 30), (30, 200, 30))):
    """造一个完整可用的 import 产物目录（不经 COM：底图用纯色 PNG 顶上）。"""
    src = _src_pptx(tmp_path, pages=len(colors))
    deck, _ = pptx_io.read_pages(src, export_width_px=640)
    out = tmp_path / "out"
    (out / "bg").mkdir(parents=True, exist_ok=True)
    for i, c in enumerate(colors, 1):
        Image.new("RGB", (64, 36), c).save(str(out / "bg" / f"slide_{i}.png"))
        deck.pages[i - 1]["bg"] = f"bg/slide_{i}.png"
    with open(out / "deck.json", "w", encoding="utf-8") as f:
        json.dump(pptx_io.deck_to_dict(deck), f, ensure_ascii=False)
    return str(out)


# ---------------------------------------------------------------- 进程契约

def test_no_subcommand_exits_2_with_json(capsys):
    code, payload, _ = _run([], capsys)
    assert code == 2
    assert payload["ok"] is False and payload["error"]["code"] == "BAD_ARGS"


def test_stdout_is_single_line_json(capsys):
    cli.main([])
    out = capsys.readouterr().out
    _assert_single_line_json(out)


def test_bad_choice_is_argparse_exit_2(capsys):
    with pytest.raises(SystemExit) as ei:
        cli.main(["import", "x.pptx", "--width", "999"])
    assert ei.value.code == 2


# ---------------------------------------------------------------- import

def test_no_com_with_faithful_is_bad_args(tmp_path, capsys):
    src = _src_pptx(tmp_path)
    code, payload, _ = _run(["import", src, "--no-com", "--mode", "faithful"], capsys)
    assert code == 2
    assert payload["error"]["code"] == "BAD_ARGS"
    assert "faithful" in payload["error"]["message"] + payload["error"]["hint"]


def test_import_missing_file(tmp_path, capsys):
    code, payload, _ = _run(["import", str(tmp_path / "nope.pptx")], capsys)
    assert code == 3 and payload["error"]["code"] == "PPTX_NOT_FOUND"


def test_import_pages_must_be_positive(tmp_path, capsys):
    src = _src_pptx(tmp_path)
    code, payload, _ = _run(["import", src, "--pages", "0"], capsys)
    assert code == 2 and payload["error"]["code"] == "BAD_ARGS"


def test_import_no_powerpoint_gives_chinese_error_not_stack(capsys, monkeypatch, tmp_path):
    """无 PowerPoint 的机器上：明确中文错误 + 可执行指引，不吐堆栈（打桩模拟）。"""
    src = _src_pptx(tmp_path)
    monkeypatch.setattr(pptx_io, "powerpoint_available", lambda: False)
    code, payload, err = _run(["import", src, "--out", str(tmp_path / "o")], capsys)
    assert code == 4
    assert payload["error"]["code"] == "NO_POWERPOINT"
    assert "PowerPoint" in payload["error"]["message"]
    assert "Traceback" not in err


def test_import_writes_deck_and_units_without_com(tmp_path, capsys, monkeypatch):
    """--mode redesign --no-com：不碰 COM 也要出 deck.json + units.json。"""
    src = _src_pptx(tmp_path, pages=2)
    monkeypatch.setattr(cli, "_design_pipeline",
                        lambda out, topic, slides, image_map=None: os.path.join(out, "deck.html"))
    import outline
    monkeypatch.setattr(outline, "generate_outline_from_text",
                        lambda text, density="balanced": [{"type": "content", "title": "t",
                                                           "points": ["a"]}])
    out = tmp_path / "o"
    code, payload, _ = _run(["import", src, "--out", str(out), "--no-com",
                             "--mode", "redesign"], capsys)
    assert code == 0, payload
    assert payload["data"]["bg"] == 0
    assert payload["data"]["pages"] == 2
    assert os.path.isfile(out / "deck.json") and os.path.isfile(out / "units.json")
    deck = pptx_io.load_deck(str(out / "deck.json"))
    assert deck.mode == "redesign" and all(p["bg"] is None for p in deck.pages)


def test_import_truncates_bg_dir_with_pages(tmp_path, capsys, monkeypatch):
    """契约 §8.2：--pages N 时 bg/ 与 deck.json 同步截断（多余底图清掉）。"""
    src = _src_pptx(tmp_path, pages=3)
    out = tmp_path / "o"

    def fake_export(pptx_path, out_dir, width=1920):
        os.makedirs(out_dir, exist_ok=True)
        made = []
        for i in range(1, 4):
            p = os.path.join(out_dir, f"slide_{i}.png")
            with open(p, "wb") as f:
                f.write(b"x")
            made.append(os.path.abspath(p))
        return made

    monkeypatch.setattr(pptx_io, "powerpoint_available", lambda: True)
    monkeypatch.setattr(pptx_io, "export_pages", fake_export)
    code, payload, _ = _run(["import", src, "--out", str(out), "--pages", "2"], capsys)
    assert code == 0, payload
    assert payload["data"]["pages"] == 2 and payload["data"]["bg"] == 2
    assert sorted(os.listdir(out / "bg")) == ["slide_1.png", "slide_2.png"]


# ---------------------------------------------------------------- animate

def test_animate_builds_player(tmp_path, capsys):
    d = _deck_dir(tmp_path)
    code, payload, _ = _run(["animate", d, "--dim", "0.3", "--auto-ms", "1500"], capsys)
    assert code == 0, payload
    player = payload["data"]["player"]
    assert os.path.isfile(player)
    doc = open(player, encoding="utf-8").read()
    assert '"dim":0.3' in doc and '"autoStepMs":1500' in doc


def test_animate_missing_deck_json(tmp_path, capsys):
    code, payload, _ = _run(["animate", str(tmp_path / "nope")], capsys)
    assert code == 3 and payload["error"]["code"] == "IR_MISMATCH"


def test_animate_missing_bg_file_is_ir_mismatch(tmp_path, capsys):
    """L4 规避：底图缺文件必须当场报错，不能让黑帧流到 video。"""
    d = _deck_dir(tmp_path)
    os.remove(os.path.join(d, "bg", "slide_2.png"))
    code, payload, _ = _run(["animate", d], capsys)
    assert code == 3 and payload["error"]["code"] == "IR_MISMATCH"
    assert "底图文件缺失" in payload["error"]["message"]


def test_animate_without_bg_is_ir_mismatch(tmp_path, capsys):
    src = _src_pptx(tmp_path, pages=1)
    deck, _ = pptx_io.read_pages(src)
    out = tmp_path / "o"
    out.mkdir()
    with open(out / "deck.json", "w", encoding="utf-8") as f:
        json.dump(pptx_io.deck_to_dict(deck), f, ensure_ascii=False)
    code, payload, _ = _run(["animate", str(out)], capsys)
    assert code == 3 and payload["error"]["code"] == "IR_MISMATCH"
    assert "faithful" in payload["error"]["hint"]


def test_animate_bad_dim_is_bad_args(tmp_path, capsys):
    d = _deck_dir(tmp_path)
    code, payload, _ = _run(["animate", d, "--dim", "0.9"], capsys)
    assert code == 2 and payload["error"]["code"] == "BAD_ARGS"


# ---------------------------------------------------------------- video

def test_video_no_ffmpeg(tmp_path, capsys, monkeypatch):
    monkeypatch.setattr(video, "ffmpeg_path", lambda: None)
    d = _deck_dir(tmp_path)
    code, payload, _ = _run(["video", d], capsys)
    assert code == 4 and payload["error"]["code"] == "NO_FFMPEG"
    assert payload["error"]["hint"]


def test_video_fade_ge_seconds_is_bad_args(tmp_path, capsys):
    """fade >= sec 会让 ffmpeg 静默产出丢页的坏视频（video.py 的既有守卫）。"""
    d = _deck_dir(tmp_path)
    code, payload, _ = _run(["video", d, "--sec", "1", "--fade", "2"], capsys)
    assert code == 2 and payload["error"]["code"] == "BAD_ARGS"


def test_video_step_without_player(tmp_path, capsys):
    d = _deck_dir(tmp_path)
    code, payload, _ = _run(["video", d, "--mode", "step"], capsys)
    assert code == 3 and payload["error"]["code"] == "IR_MISMATCH"


# ---------------------------------------------------------------- export

def test_export_missing_deck_json(tmp_path, capsys):
    code, payload, _ = _run(["export", str(tmp_path / "nope.json")], capsys)
    assert code == 3 and payload["error"]["code"] == "PPTX_NOT_FOUND"


def test_export_bad_template(tmp_path, capsys):
    d = _deck_dir(tmp_path)
    bad = tmp_path / "brand.txt"
    bad.write_text("not a pptx", encoding="utf-8")
    code, payload, _ = _run(["export", os.path.join(d, "deck.json"),
                             "--template", str(bad)], capsys)
    assert code == 3 and payload["error"]["code"] == "TEMPLATE_INVALID"


# ---------------------------------------------------------------- deck

def test_deck_command_runs_pipeline(tmp_path, capsys, monkeypatch):
    import outline
    monkeypatch.setattr(outline, "generate_outline",
                        lambda topic, density="balanced": [{"type": "cover", "title": topic}])
    monkeypatch.setattr(cli, "_design_pipeline",
                        lambda out, topic, slides, image_map=None: os.path.join(out, "deck.html"))
    out = tmp_path / "o"
    code, payload, _ = _run(["deck", "人工智能", "--out", str(out)], capsys)
    assert code == 0, payload
    assert payload["data"]["pages"] == 1 and payload["data"]["topic"] == "人工智能"


# ---------------------------------------------------------------- 底图路径基准

def test_rebase_bg_rewrites_to_player_dir(tmp_path):
    """契约 §5 的 <out>/player/ 布局：src 必须是 ../bg/x.png 而不是 bg/x.png。

    这是真实踩到的 bug —— animate 出播放器后底图 3/3 全部加载失败（页面全黑），
    靠 shot_player 的 L4 校验才暴露出来。
    """
    out_dir = str(tmp_path / "player")
    got = hl_anim._rebase_bg("bg/slide_1.png", out_dir, str(tmp_path))
    assert got == "../bg/slide_1.png"


def test_rebase_bg_rejects_escape(tmp_path):
    from pptx_io import PptxError
    with pytest.raises(PptxError) as ei:
        hl_anim._rebase_bg("../outside.png", str(tmp_path / "player"), str(tmp_path))
    assert ei.value.code == "IR_MISMATCH"


def test_build_player_with_bg_base_dir(tmp_path):
    """bg_base_dir 参与后，播放器里的 src 指向真实的底图位置。"""
    base = tmp_path / "out"
    (base / "bg").mkdir(parents=True)
    Image.new("RGB", (8, 8), (1, 2, 3)).save(str(base / "bg" / "slide_1.png"))
    player = hl_anim.build_player(str(base / "player"), ["bg/slide_1.png"], [[]],
                                  bg_base_dir=str(base))
    doc = open(player, encoding="utf-8").read()
    assert "../bg/slide_1.png" in doc
    # 浏览器解析后的绝对路径确实存在
    resolved = os.path.normpath(os.path.join(os.path.dirname(player), "../bg/slide_1.png"))
    assert os.path.isfile(resolved)


# ---------------------------------------------------------------- 端到端（真相机）

_needs_browser = pytest.mark.skipif(
    subprocess.run([sys.executable, "-c", "import playwright"],
                   capture_output=True).returncode != 0,
    reason="未安装 playwright")


@_needs_browser
def test_animate_then_video_step_end_to_end(tmp_path, capsys):
    """animate → video --mode step 全链路：帧数 == Σ units，产物非空。"""
    d = _deck_dir(tmp_path)
    code, payload, _ = _run(["animate", d], capsys)
    assert code == 0, payload
    code, payload, _ = _run(["video", d, "--mode", "step", "--sec", "1",
                             "--fade", "0.4", "-o", str(tmp_path / "s.mp4")], capsys)
    if code == 4 and payload["error"]["code"] == "NO_BROWSER":
        pytest.skip("无可用浏览器")
    assert code == 0, payload
    assert payload["data"]["clips"] == 2  # 2 页 × 各 1 个单元
    assert os.path.getsize(payload["data"]["video"]) > 0
