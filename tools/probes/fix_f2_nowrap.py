"""F2 对照探针：`wrap="none"` 的文本框让"行级高亮"退化成"段级"，但没人告诉你。

背景（TEST_REPORT §6 F2）：python-pptx 的 `shapes.add_textbox()` **默认写 `wrap="none"`**
（`output/b_multislide.pptx` 里 13/22 个文本框如此）。这类框我们按"不换行"处理
（契约 §4.3 规定），于是一个段落恒出 1 行 —— 行级与段级在那里是同一件事。
如果文本估算宽度还**超过**框宽，这块高亮要么真的溢出、要么我们无从知道它会怎么折
（**不猜**是契约的选择），但调用方至少该被告知"这里的精度是段级"。

判定：
  修复前 → Unit.warnings 里没有降级标记；CLI warnings 也是空的
  修复后 → 出现 `nowrap_overflow_degraded`；CLI warnings 能看到计数

跑法：.venv/Scripts/python.exe tools/probes/fix_f2_nowrap.py
"""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, ROOT)

SRC = os.path.join(ROOT, "output", "b_multislide.pptx")
MARK = "nowrap_overflow_degraded"


def main() -> int:
    if not os.path.isfile(SRC):
        print(f"SKIP: 缺少素材 {SRC}")
        return 0

    import hl_layout
    import pptx_io

    deck, _ = pptx_io.read_pages(SRC)
    print(f"素材：{SRC}")

    wrap_none = overflow = 1
    total_units = 0
    for page in deck.pages:
        for sh in page["shapes"]:
            if sh.kind != "text":
                continue
            if sh.word_wrap is False:
                wrap_none += 1
    print(f"  word_wrap=False 的文本框：{wrap_none - 1} 个（python-pptx 的默认值）")

    marked_units = []
    total_units = 0
    for page in deck.pages:
        units = hl_layout.build_units(hl_layout.page_shapes(page, deck))
        for u in units:
            total_units += 1
            if MARK in u.warnings:
                marked_units.append(u)

    print(f"\n== build_units 的 Unit.warnings ==")
    print(f"  讲解单元合计 = {total_units}")
    print(f"  带 `{MARK}` 标记的单元 = {len(marked_units)}")
    for u in marked_units[:5]:
        est_w = max(r.width_px for r in u.lines)
        print(f"    · 第{u.page_index + 1}页 {u.shape_name}: {u.size_pt:.0f}pt "
              f"单行估宽 {est_w:.0f}px，文本 {u.text[:14]!r}")

    print(f"\n== 走 CLI：warnings 里能不能看到 ==")
    import io
    import json
    from contextlib import redirect_stdout

    import cli
    import outline
    outline.generate_outline_from_text = lambda text, density="balanced": []
    cli._design_pipeline = lambda out, topic, slides, image_map=None: ""

    buf = io.StringIO()
    out_dir = os.path.join(ROOT, "output", "spike", "f2probe")
    with redirect_stdout(buf):
        rc = cli.main(["import", SRC, "--no-com", "--mode", "redesign",
                       "--out", os.path.abspath(out_dir)])
    payload = json.loads(buf.getvalue())
    warnings = payload.get("data", {}).get("warnings", [])
    print(f"  CLI 退出码 = {rc}")
    print(f"  warnings = {warnings}")

    ok = marked_units and any(MARK in w for w in warnings)
    print("\n" + "-" * 68)
    if ok:
        print("判定：**FIXED** —— 段级降级有出口，Unit 与 CLI 都能看到")
    else:
        print("判定（b_multislide）：**没有单元命中** —— 见下方阳性对照")
    print("-" * 68)

    _positive_control(tmp=None)
    return 0


def _positive_control(tmp=None) -> int:
    """阳性对照：造一个 `wrap="none"` 且文本明显超出框宽的形状，警告必须出现。

    为什么需要它：`b_multislide.pptx` 里 13 个 `wrap="none"` 框的文本**都比框窄**
    （实测：内宽 453~830pt，首段估宽 56~528pt），所以"估宽 > 框宽"这个条件在该稿上
    **一处都不成立** —— 任务卡要求的"对该稿跑 import 应在 warnings 里看到降级提示"
    因此无法触发。用合成形状证明判据本身是通的。
    """
    import shutil
    import tempfile

    from pptx import Presentation
    from pptx.util import Inches, Pt

    print("\n\n== 阳性对照：wrap='none' 且文本超出框宽 ==")
    tmp = tempfile.mkdtemp(prefix="f2ctl_")
    try:
        src = os.path.join(tmp, "ctl.pptx")
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        box.name = "NarrowNowrap"
        box.text_frame.word_wrap = False          # python-pptx 的默认值，这里显式写死
        box.text_frame.text = "这一行文字明显比这个两英寸宽的文本框要长很多很多"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        prs.save(src)

        import hl_layout
        import pptx_io
        deck, _ = pptx_io.read_pages(src)
        units = hl_layout.build_units(hl_layout.page_shapes(deck.pages[0], deck))
        marks = [w for u in units for w in u.warnings if w == MARK]
        print(f"  单元数 = {len(units)}，带 `{MARK}` 的警告数 = {len(marks)}")
        ok = len(marks) == 1
        print(f"  判定：{'PASS（判据本身是通的）' if ok else 'FAIL（判据没生效）'}")
        return 0 if ok else 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
