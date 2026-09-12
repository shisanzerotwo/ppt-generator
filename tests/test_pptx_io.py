"""M1：pptx_io.read_pages / export_pages / PptxError 单测。

分两层：
- 纯读层（不需要 PowerPoint）：所有机器都能跑；
- COM 层：无 PowerPoint 则 skip（照抄 test_shot_settle 的 skipif 模式）。
"""

import os

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import pptx_io

SRC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "output", "b_multislide.pptx")

needs_pptx_src = pytest.mark.skipif(not os.path.isfile(SRC),
                                    reason="缺 output/b_multislide.pptx 素材")


# ---------------------------------------------------------------- 造稿工具

def _blank_deck(path, pages=1):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for _ in range(pages):
        prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))
    return str(path)


def _box(slide, text, size=18.0, name=None):
    box = slide.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(4.0), Inches(1.5))
    if name:
        box.name = name
    box.text_frame.word_wrap = True
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    return box


# ---------------------------------------------------------------- 读形状

@needs_pptx_src
def test_reads_ten_pages_with_coords():
    """M1 验收：10 页、坐标非 None（原始 41 形状 = 40 文本框 + 1 图表）。"""
    deck, skipped = pptx_io.read_pages(SRC)
    assert len(deck.pages) == 10
    assert deck.width_emu == 12191695 and deck.height_emu == 6858000
    assert deck.schema == 1 and deck.mode == "faithful"

    kept = [s for p in deck.pages for s in p["shapes"]]
    assert {s.kind for s in kept} == {"text", "chart"}
    for s in kept:
        assert None not in (s.left_emu, s.top_emu, s.width_emu, s.height_emu)
        assert s.width_emu > 0 and s.height_emu > 0


@needs_pptx_src
def test_empty_textboxes_are_filtered():
    """spike 口径：40 个文本框里 18 个是空的，导入期一次性过滤。"""
    _, skipped = pptx_io.read_pages(SRC)
    empty = [x for x in skipped if x["reason"] == "empty_text"]
    assert len(empty) == 18
    assert all(x["page_index"] >= 0 for x in empty)


def test_page_bg_defaults_to_none(tmp_path):
    """bg 由 export_pages 之后回填 —— read_pages 不知道 out 目录。"""
    deck, _ = pptx_io.read_pages(_blank_deck(tmp_path / "a.pptx"))
    assert all(p["bg"] is None for p in deck.pages)


def test_manual_line_break_becomes_newline(tmp_path):
    """段内 a:br 必须读成 "\\n"：paragraph.runs 不暴露它，直接用会丢换行。"""
    path = tmp_path / "br.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _box(slide, "第一行")
    p = box.text_frame.paragraphs[0]
    p.runs[0].text = "第一行"
    br = p._p.makeelement(qn("a:br"), {})
    p._p.append(br)
    r2 = p._p.makeelement(qn("a:r"), {})
    t2 = p._p.makeelement(qn("a:t"), {})
    t2.text = "第二行"
    r2.append(t2)
    p._p.append(r2)
    prs.save(str(path))

    deck, _ = pptx_io.read_pages(str(path))
    para = deck.pages[0]["shapes"][0].paragraphs[0]
    assert para.text == "第一行\n第二行"


def test_run_level_font_fields_are_preserved(tmp_path):
    """size/bold/italic/latin 逐 run 读；取不到就保留 None（回退留给 hl_layout）。"""
    path = tmp_path / "runs.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    p = box.text_frame.paragraphs[0]
    r1 = p.add_run(); r1.text = "大"; r1.font.size = Pt(32); r1.font.bold = True
    r2 = p.add_run(); r2.text = "小"
    prs.save(str(path))

    deck, _ = pptx_io.read_pages(str(path))
    runs = deck.pages[0]["shapes"][0].paragraphs[0].runs
    assert (runs[0].text, runs[0].size_pt, runs[0].bold) == ("大", 32.0, True)
    assert (runs[1].text, runs[1].size_pt) == ("小", None)


def test_line_spacing_multiple_and_length(tmp_path):
    """倍数原样存；Length 形态按契约折算为"倍数"再存。"""
    path = tmp_path / "ls.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(3))
    tf = box.text_frame
    tf.text = "倍数行距"
    tf.paragraphs[0].runs[0].font.size = Pt(16)
    tf.paragraphs[0].line_spacing = 1.8
    p2 = tf.add_paragraph(); p2.add_run().text = "固定行距"
    p2.runs[0].font.size = Pt(20)
    p2.line_spacing = Pt(30)
    prs.save(str(path))

    deck, _ = pptx_io.read_pages(str(path))
    paras = deck.pages[0]["shapes"][0].paragraphs
    assert paras[0].line_spacing == pytest.approx(1.8)
    assert paras[1].line_spacing == pytest.approx(30 / 20)


def test_space_before_after_and_align(tmp_path):
    path = tmp_path / "sb.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(3))
    tf = box.text_frame
    tf.text = "标题"
    tf.paragraphs[0].space_before = Pt(20)
    tf.paragraphs[0].space_after = Pt(8)
    from pptx.enum.text import PP_ALIGN
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    prs.save(str(path))

    deck, _ = pptx_io.read_pages(str(path))
    p = deck.pages[0]["shapes"][0].paragraphs[0]
    assert p.space_before_pt == 20.0 and p.space_after_pt == 8.0
    assert p.align == "CENTER" and p.level == 0


def test_font_scale_read_from_norm_autofit(tmp_path):
    """a:normAutofit/@fontScale ÷ 1000；没有该元素则为 None。"""
    path = tmp_path / "fs.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _box(slide, "会被自动缩排的文字")
    bodyPr = box.text_frame._txBody.find(qn("a:bodyPr"))
    bodyPr.append(bodyPr.makeelement(qn("a:normAutofit"), {"fontScale": "75000"}))
    prs.save(str(path))

    deck, _ = pptx_io.read_pages(str(path))
    assert deck.pages[0]["shapes"][0].font_scale == pytest.approx(75.0)


def test_text_frame_attributes_read(tmp_path):
    """word_wrap / vertical_anchor / auto_size 原样读，None 保留 None（= 继承）。"""
    path = tmp_path / "attrs.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _box(slide, "内容")
    from pptx.enum.text import MSO_ANCHOR
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    prs.save(str(path))

    deck, _ = pptx_io.read_pages(str(path))
    sh = deck.pages[0]["shapes"][0]
    assert sh.vertical_anchor == "MIDDLE"
    assert sh.margin_left_emu == 91440 and sh.margin_top_emu == 45720


# ---------------------------------------------------------------- 过滤规则

def test_hidden_shape_is_skipped(tmp_path):
    path = tmp_path / "hidden.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _box(slide, "藏起来", name="Hider")
    box._element.nvSpPr.cNvPr.set("hidden", "1")
    prs.save(str(path))

    deck, skipped = pptx_io.read_pages(str(path))
    assert deck.pages[0]["shapes"] == []
    assert skipped == [{"page_index": 0, "shape_name": "Hider", "reason": "hidden"}]


def test_zero_size_shape_is_skipped(tmp_path):
    path = tmp_path / "zero.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), 0, Inches(1))
    box.text_frame.text = "零宽"
    prs.save(str(path))

    _, skipped = pptx_io.read_pages(str(path))
    assert [x["reason"] for x in skipped] == ["zero_size"]


def test_unsupported_shape_kept_for_redesign(tmp_path):
    """kind=="other" 保号留在 shapes 里（供 redesign），但记 unsupported 不进 units。"""
    path = tmp_path / "other.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_connector(1, Inches(1), Inches(1), Inches(3), Inches(3))
    prs.save(str(path))

    deck, skipped = pptx_io.read_pages(str(path))
    kinds = [s.kind for s in deck.pages[0]["shapes"]]
    assert kinds == ["other"]
    assert [x["reason"] for x in skipped] == ["unsupported"]


# ---------------------------------------------------------------- 组合形状

def _group_deck(path, scaled=True):
    """组内两个文本框；chExt 比 ext 小一半 → 触发 2x 仿射放大。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.name = "G"
    t1 = group.shapes.add_textbox(Emu(0), Emu(0), Emu(1000000), Emu(500000))
    t1.text_frame.text = "组内一"
    t2 = group.shapes.add_textbox(Emu(0), Emu(1000000), Emu(1000000), Emu(500000))
    t2.text_frame.text = "组内二"
    xfrm = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    off.set("x", str(Inches(2))); off.set("y", str(Inches(2)))
    ext.set("cx", str(Inches(4))); ext.set("cy", str(Inches(2)))
    ch_off.set("x", "0"); ch_off.set("y", "0")
    if scaled:
        ch_ext.set("cx", str(Inches(2))); ch_ext.set("cy", str(Inches(2)))
    else:
        ch_ext.set("cx", str(Inches(4))); ch_ext.set("cy", str(Inches(2)))
    prs.save(str(path))
    return str(path)


def test_group_children_are_affine_mapped(tmp_path):
    """契约 §2.2：子坐标在子坐标系里，必须按 off/ext 与 chOff/chExt 做仿射。"""
    path = _group_deck(tmp_path / "grp.pptx", scaled=True)
    deck, skipped = pptx_io.read_pages(str(path))
    grp = deck.pages[0]["shapes"][0]
    assert grp.kind == "group" and len(grp.children) == 2

    gx = int(Inches(2))
    # chOff=(0,0) chExt=(2in,2in) → 水平放大 2 倍；子框宽 1000000 EMU → 2000000
    assert grp.children[0].left_emu == gx
    assert grp.children[0].width_emu == pytest.approx(2000000, abs=1)
    assert grp.children[1].top_emu == pytest.approx(int(Inches(2)) + 1000000, abs=1)
    assert not [x for x in skipped if x["reason"] == "degenerate_group"]


def test_group_identity_when_ch_ext_equals_ext(tmp_path):
    """chExt == ext → 恒等映射，子坐标原样落到页面上。"""
    path = _group_deck(tmp_path / "grp2.pptx", scaled=False)
    deck, _ = pptx_io.read_pages(str(path))
    child = deck.pages[0]["shapes"][0].children[1]
    assert child.left_emu == int(Inches(2))
    assert child.top_emu == int(Inches(2)) + 1000000


def test_group_without_xfrm_is_skipped_whole(tmp_path):
    """组自身没有 a:xfrm → 四个属性全 None，整组跳过（实测行为）。"""
    path = _group_deck(tmp_path / "grp3.pptx")
    prs = Presentation(path)
    grp = next(s for s in prs.slides[0].shapes if s.shape_type == 6)
    grpSpPr = grp._element.find(qn("p:grpSpPr"))
    grpSpPr.remove(grpSpPr.find(qn("a:xfrm")))
    prs.save(path)

    deck, skipped = pptx_io.read_pages(path)
    assert deck.pages[0]["shapes"] == []
    assert [x["reason"] for x in skipped] == ["group_no_xfrm"]


def test_nested_group_mapping_is_cumulative(tmp_path):
    """嵌套组合逐层累乘，不能只对顶层算一次。"""
    path = tmp_path / "nested.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_textbox(Emu(0), Emu(0), Emu(100000), Emu(100000)).text_frame.text = "深处"
    for grp, ox, oy, gw, gh, cw, ch in ((outer, Inches(1), Inches(1), Inches(4), Inches(4),
                                         Inches(2), Inches(2)),
                                        (inner, 0, 0, Inches(2), Inches(2),
                                         Inches(1), Inches(1))):
        xfrm = grp._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
        xfrm.find(qn("a:off")).set("x", str(int(ox)))
        xfrm.find(qn("a:off")).set("y", str(int(oy)))
        xfrm.find(qn("a:ext")).set("cx", str(int(gw)))
        xfrm.find(qn("a:ext")).set("cy", str(int(gh)))
        xfrm.find(qn("a:chOff")).set("x", "0")
        xfrm.find(qn("a:chOff")).set("y", "0")
        xfrm.find(qn("a:chExt")).set("cx", str(int(cw)))
        xfrm.find(qn("a:chExt")).set("cy", str(int(ch)))
    prs.save(str(path))

    deck, _ = pptx_io.read_pages(str(path))
    outer_sh = deck.pages[0]["shapes"][0]
    inner_sh = outer_sh.children[0]
    leaf = inner_sh.children[0]
    # 外层 1in + 内层 0 → 1in；宽度 100000 EMU 经 2x 与 2x 两次放大 = 400000
    assert leaf.left_emu == int(Inches(1))
    assert leaf.width_emu == pytest.approx(100000 * 4, abs=1)


# ---------------------------------------------------------------- 表格

def test_merged_table_cells_are_read(tmp_path):
    path = tmp_path / "tbl.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(9), Inches(3))
    tbl = shape.table
    for r in range(2):
        for c in range(3):
            tbl.cell(r, c).text_frame.text = f"r{r}c{c}"
    tbl.cell(1, 0).merge(tbl.cell(1, 2))
    prs.save(str(path))

    deck, _ = pptx_io.read_pages(str(path))
    sh = deck.pages[0]["shapes"][0]
    assert sh.kind == "table"
    assert len(sh.table_col_widths_emu) == 3 and len(sh.table_row_heights_emu) == 2
    assert sh.table_cells[0][1][0].text == "r0c1"

    prs2 = Presentation(str(path))
    t2 = next(s for s in prs2.slides[0].shapes if s.has_table).table
    assert t2.cell(1, 0).is_merge_origin and t2.cell(1, 0).span_width == 3
    assert t2.cell(1, 1).is_spanned and t2.cell(1, 2).is_spanned


# ---------------------------------------------------------------- 错误码

def test_not_found_code(tmp_path):
    with pytest.raises(pptx_io.PptxError) as ei:
        pptx_io.read_pages(str(tmp_path / "nope.pptx"))
    assert ei.value.code == "PPTX_NOT_FOUND"
    assert ei.value.to_dict()["hint"]


def test_unreadable_code(tmp_path):
    p = tmp_path / "bad.pptx"
    p.write_bytes(b"this is definitely not a zip")
    with pytest.raises(pptx_io.PptxError) as ei:
        pptx_io.read_pages(str(p))
    assert ei.value.code == "PPTX_UNREADABLE"


def test_encrypted_code(tmp_path):
    """加密的 Office 文件是 OLE2 复合文档（D0CF11E0...），不是 zip。"""
    p = tmp_path / "enc.pptx"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512)
    with pytest.raises(pptx_io.PptxError) as ei:
        pptx_io.read_pages(str(p))
    assert ei.value.code == "PPTX_ENCRYPTED"


def test_empty_deck_code(tmp_path):
    p = tmp_path / "empty.pptx"
    prs = Presentation()  # 默认模板 0 页
    prs.save(str(p))
    with pytest.raises(pptx_io.PptxError) as ei:
        pptx_io.read_pages(str(p))
    assert ei.value.code == "PPTX_EMPTY"


def test_pptx_error_to_dict_shape():
    err = pptx_io.PptxError("消息", "CODE", "提示")
    assert err.to_dict() == {"code": "CODE", "message": "消息", "hint": "提示"}
    assert str(err) == "消息"


# ---------------------------------------------------------------- deck.json

@needs_pptx_src
def test_deck_round_trip(tmp_path):
    import json
    deck, _ = pptx_io.read_pages(SRC)
    d = pptx_io.deck_to_dict(deck)
    path = tmp_path / "deck.json"
    path.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")

    back = pptx_io.load_deck(str(path))
    assert back.source_pptx == deck.source_pptx
    assert len(back.pages) == len(deck.pages)
    assert back.pages[0]["shapes"][0].paragraphs[0].text == \
        deck.pages[0]["shapes"][0].paragraphs[0].text


def test_load_deck_bad_json(tmp_path):
    p = tmp_path / "deck.json"
    p.write_text("{not json", encoding="utf-8")
    with pytest.raises(pptx_io.PptxError) as ei:
        pptx_io.load_deck(str(p))
    assert ei.value.code == "IR_MISMATCH"


def test_load_deck_missing_key(tmp_path):
    p = tmp_path / "deck.json"
    p.write_text('{"pages": []}', encoding="utf-8")
    with pytest.raises(pptx_io.PptxError) as ei:
        pptx_io.load_deck(str(p))
    assert ei.value.code == "IR_MISMATCH"


# ---------------------------------------------------------------- COM 导出

_no_ppt = pytest.mark.skipif(not pptx_io.powerpoint_available(),
                             reason="本机没有 PowerPoint（COM 导出不可用）")


def test_powerpoint_available_returns_bool():
    assert isinstance(pptx_io.powerpoint_available(), bool)


@_no_ppt
def test_export_pages_size_and_order(tmp_path):
    """导出 3 页：张数/页序/像素尺寸对齐，文件名 slide_1.png 起。"""
    src = _blank_deck(tmp_path / "src.pptx", pages=3)
    out = tmp_path / "bg"
    paths = pptx_io.export_pages(src, str(out), width=1920)

    assert len(paths) == 3
    assert [os.path.basename(p) for p in paths] == ["slide_1.png", "slide_2.png", "slide_3.png"]
    from PIL import Image
    for p in paths:
        assert os.path.isabs(p) and os.path.getsize(p) > 0
        with Image.open(p) as im:
            assert im.size == (1920, 1080)


@_no_ppt
def test_export_pages_width_scales_height(tmp_path):
    """1280 宽 → 高按画布比例取整为 720。"""
    src = _blank_deck(tmp_path / "src.pptx", pages=1)
    paths = pptx_io.export_pages(src, str(tmp_path / "bg"), width=1280)
    from PIL import Image
    with Image.open(paths[0]) as im:
        assert im.size == (1280, 720)


@_no_ppt
def test_export_pages_missing_file(tmp_path):
    with pytest.raises(pptx_io.PptxError) as ei:
        pptx_io.export_pages(str(tmp_path / "nope.pptx"), str(tmp_path / "o"))
    assert ei.value.code == "PPTX_NOT_FOUND"


def test_export_pages_reports_no_powerpoint_without_com(monkeypatch, tmp_path):
    """无 PowerPoint 的机器上要给明确中文错误，不是裸堆栈（打桩模拟）。"""
    monkeypatch.setattr(pptx_io, "powerpoint_available", lambda: False)
    src = _blank_deck(tmp_path / "src.pptx", pages=1)
    with pytest.raises(pptx_io.PptxError) as ei:
        pptx_io.export_pages(src, str(tmp_path / "bg"))
    assert ei.value.code == "NO_POWERPOINT"
    assert "PowerPoint" in ei.value.message and ei.value.hint


def test_export_pages_rejects_bad_width(monkeypatch, tmp_path):
    monkeypatch.setattr(pptx_io, "powerpoint_available", lambda: True)
    src = _blank_deck(tmp_path / "src.pptx", pages=1)
    with pytest.raises(pptx_io.PptxError) as ei:
        pptx_io.export_pages(src, str(tmp_path / "bg"), width=0)
    assert ei.value.code == "BAD_ARGS"
