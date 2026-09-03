"""builder 模块纯逻辑 + build_ppt 离线集成测试。"""

import os

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from builder import (
    _fit_title_size,
    _fit_body_size,
    _pick_chart_type,
    build_ppt,
    THEMES,
)


# ---------- _fit_title_size ----------

def test_fit_title_size_short():
    assert _fit_title_size("短标题") == 28
    assert _fit_title_size("1" * 14) == 28  # 边界：14 字仍 28


def test_fit_title_size_medium():
    assert _fit_title_size("1" * 15) == 24  # 15 字
    assert _fit_title_size("1" * 20) == 24  # 边界：20 字


def test_fit_title_size_long():
    assert _fit_title_size("1" * 21) == 20  # 21 字
    assert _fit_title_size("1" * 28) == 20  # 边界：28 字


def test_fit_title_size_very_long():
    assert _fit_title_size("1" * 29) == 18
    assert _fit_title_size("一" * 29) == 18


# ---------- _fit_body_size ----------

def test_fit_body_size_small():
    assert _fit_body_size(["a" * 60]) == 17  # 总长 60
    assert _fit_body_size(["短", "更短"]) == 17


def test_fit_body_size_medium():
    assert _fit_body_size(["a" * 61]) == 15  # 总长 61
    assert _fit_body_size(["a" * 110]) == 15  # 边界：110


def test_fit_body_size_large():
    assert _fit_body_size(["a" * 111]) == 13
    assert _fit_body_size(["a" * 60, "b" * 60]) == 13  # 总长 120


def test_fit_body_size_empty():
    assert _fit_body_size([]) == 17  # 总长 0


# ---------- _pick_chart_type ----------

def test_pick_chart_type_explicit_type_wins():
    # labels 像时间序列本应选 line，但显式 column 优先
    assert _pick_chart_type({"type": "column", "labels": ["2015", "2016"], "values": [10, 20]}) == "column"
    assert _pick_chart_type({"type": "bar", "labels": ["A"], "values": [1]}) == "bar"


def test_pick_chart_type_ignores_unknown_explicit_type():
    # 未知 type 忽略，走自动选型（labels 含数字 → line）
    assert _pick_chart_type({"type": "area", "labels": ["2015"], "values": [10]}) == "line"


def test_pick_chart_type_time_series_to_line():
    assert _pick_chart_type({"labels": ["2015年", "2016年", "2017年"], "values": [10, 20, 30]}) == "line"


def test_pick_chart_type_share_to_pie():
    assert _pick_chart_type({"labels": ["A", "B"], "values": [60, 40]}) == "pie"  # 总和 100


def test_pick_chart_type_default_column():
    # labels 非时间序列，values 总和偏离 100 → column
    assert _pick_chart_type({"labels": ["甲", "乙"], "values": [10, 20]}) == "column"


def test_pick_chart_type_empty_values_default_column():
    assert _pick_chart_type({"labels": ["A"], "values": []}) == "column"


# ---------- build_ppt ----------

def _sample_slides():
    return [
        {"type": "cover", "title": "主题", "points": [], "image_prompt": "封面插画"},
        {"type": "toc", "title": "目录", "points": ["章节一", "章节二"], "image_prompt": ""},
        {"type": "section", "title": "第一章", "points": ["一句话简介"], "image_prompt": ""},
        {"type": "content", "title": "要点页", "points": ["要点1", "要点2", "要点3"], "image_prompt": "场景图"},
        {"type": "data", "title": "数据页", "points": [], "image_prompt": "",
         "chart": {"type": "column", "labels": ["A", "B"], "values": [30, 70]}},
        {"type": "timeline", "title": "发展历程", "points": ["2015年：成立", "2020年：扩张"], "image_prompt": ""},
        {"type": "compare", "title": "对比", "points": ["左1", "左2", "右1", "右2"], "image_prompt": ""},
        {"type": "end", "title": "感谢观看", "points": [], "image_prompt": ""},
    ]


def test_build_ppt_all_eight_types(tmp_path):
    slides = _sample_slides()
    out_path = str(tmp_path / "out.pptx")
    result = build_ppt(slides, [None] * len(slides), out_path, theme="blue", subtitle="副标题")

    assert result == out_path
    assert os.path.exists(out_path)

    prs = Presentation(out_path)
    assert len(prs.slides) == 8

    # data 页应含图表形状，且类型为 column
    chart_shapes = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(chart_shapes) == 1
    assert chart_shapes[0].chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED


def test_build_ppt_pie_chart(tmp_path):
    slides = [
        {"type": "cover", "title": "T", "points": [], "image_prompt": "p"},
        {"type": "data", "title": "D", "points": [], "image_prompt": "",
         "chart": {"type": "pie", "labels": ["A", "B"], "values": [60, 40]}},
    ]
    out_path = str(tmp_path / "pie.pptx")
    build_ppt(slides, [None, None], out_path)

    prs = Presentation(out_path)
    assert len(prs.slides) == 2
    chart_shapes = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(chart_shapes) == 1
    assert chart_shapes[0].chart.chart_type == XL_CHART_TYPE.PIE


def test_build_ppt_unknown_theme_falls_back_to_blue(tmp_path):
    slides = [{"type": "cover", "title": "T", "points": [], "image_prompt": "p"}]
    out_path = str(tmp_path / "theme.pptx")
    build_ppt(slides, [None], out_path, theme="不存在的主题")
    assert os.path.exists(out_path)


def test_themes_contain_expected_keys():
    assert "blue" in THEMES and "dark" in THEMES and "green" in THEMES
    for t in THEMES.values():
        assert {"bg", "accent", "fg", "muted"} <= set(t.keys())
