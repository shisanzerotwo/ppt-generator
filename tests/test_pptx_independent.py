"""独立验证用例集（ppt-test 角色）· 与实现者自测**不同方法/不同素材**的交叉检查。

只新增本文件，不改任何实现与既有测试。包含：
  1. 覆盖率度量的**独立像素法**交叉验证（含 ground-truth 对照图，锁住已知盲点）；
  2. `wrap_lines` 与 `qa.measure_text_lines` 的等价性 + 三条独立不变量；
  3. 播放器转义对抗用例（含真实浏览器，无浏览器则 skip）；
  4. 输入病态与边界路径的错误码；
  5. M1 口径：裸 XML 数形状（依赖 output/b_multislide.pptx，缺失则 skip）。

探针版（含更重的真机 COM 检查）在 `tools/probes/indep_*.py`。
"""

import os
import re
import zipfile
from collections import Counter

import pytest

import hl_anim
import hl_layout
import pptx_io
import qa
from hl_layout import Rect, Unit
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
B_MULTISLIDE = os.path.join(PROJECT_ROOT, "output", "b_multislide.pptx")
BG_DIR = os.path.join(PROJECT_ROOT, "output", "spike", "m1")


# ================================================================ 1. 独立像素法

def _png(path, size, bg=(255, 255, 255)):
    from PIL import Image
    img = Image.new("RGB", size, bg)
    img.save(path)
    return str(path), img


def _far_ring_coverage(path, box, thresh=40, far=8):
    """独立口径：底色取 rect **外** far 像素环的众数色，墨迹判据用 RGB 欧氏距离。"""
    from PIL import Image
    img = Image.open(path).convert("RGB")
    W, H = img.size
    l, t, w, h = box
    r, b = l + w, t + h
    px = img.load()
    coords = []
    for band in range(far, far + 3):
        for x in range(max(0, l - band), min(W, r + band)):
            for y in (t - band, b + band - 1):
                if 0 <= y < H:
                    coords.append((x, y))
        for y in range(max(0, t - band), min(H, b + band)):
            for x in (l - band, r + band - 1):
                if 0 <= x < W:
                    coords.append((x, y))
    counts = {}
    for x, y in coords:
        p = px[x, y]
        counts[(p[0] // 8, p[1] // 8, p[2] // 8)] = counts.get((p[0] // 8, p[1] // 8, p[2] // 8), 0) + 1
    k = max(counts, key=counts.get)
    bg = (k[0] * 8 + 4, k[1] * 8 + 4, k[2] * 8 + 4)
    ink = 0
    for y in range(max(0, t), min(H, b)):
        for x in range(max(0, l), min(W, r)):
            p = px[x, y]
            dr, dg, db = p[0] - bg[0], p[1] - bg[1], p[2] - bg[2]
            if dr * dr + dg * dg + db * db > thresh * thresh:
                ink += 1
    return ink / (w * h)


def test_independent_method_matches_ground_truth(tmp_path):
    """对照图（已知 30% 墨迹）：独立法与实现者的 measure_coverage 都必须测准。"""
    from PIL import Image, ImageDraw
    p = str(tmp_path / "gt30.png")
    img = Image.new("RGB", (200, 100), (255, 255, 255))
    ImageDraw.Draw(img).rectangle([0, 0, 99, 29], fill=(0, 0, 0))
    img.save(p)
    box = (0, 0, 100, 100)
    assert hl_layout.measure_coverage(p, box) == pytest.approx(0.30, abs=0.02)
    assert _far_ring_coverage(p, box) == pytest.approx(0.30, abs=0.02)


def test_contract_default_bg_is_blind_when_ink_fills_rect_border(tmp_path):
    """锁住 contract §4.5 默认口径（内侧 1px 环取中位色）的**盲点**：

    rect 恰好等于墨迹块时，环本身就是墨迹色 → 底色被判成墨迹色 → 覆盖率 0.0，
    而独立法（环取 rect **外**）得到真实值 1.0。本素材上环上墨迹占比 ≤0.04 所以
    不发作（见 tools/probes/indep_m2_coverage.py 第 3 节），但口径本身有这个空洞。
    """
    from PIL import Image, ImageDraw
    p = str(tmp_path / "border_ink.png")
    img = Image.new("RGB", (200, 100), (255, 255, 255))
    ImageDraw.Draw(img).rectangle([20, 20, 59, 59], fill=(0, 0, 0))
    img.save(p)
    box = (20, 20, 40, 40)
    assert hl_layout.measure_coverage(p, box) == 0.0          # 契约默认口径：测不出
    assert _far_ring_coverage(p, box) == pytest.approx(1.0, abs=0.02)   # 独立口径：正确


def test_fill_colored_pixels_inflate_coverage_with_wrong_bg(tmp_path):
    """口径差异（impl 报告 D2 的 0.391 vs 0.286）：rect 落在浅色填充形状上时，

    - 契约默认（内侧环取中位色）→ 底色 = 填充色 → 只把**文字**算墨迹（正确）；
    - 显式传"幻灯片主色"（白）→ 填充色整块被算成墨迹 → 覆盖率**虚高**。
    两法都保留为断言，锁住"验收脚本必须显式传局部底色"这条经验。
    """
    from PIL import Image, ImageDraw
    p = str(tmp_path / "on_fill.png")
    img = Image.new("RGB", (200, 120), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.rectangle([20, 20, 179, 99], fill=(182, 199, 220))     # 浅色填充形状
    for y in (40, 56, 72):                                   # 三行"文字"
        d.rectangle([40, y, 139, y + 7], fill=(20, 20, 20))
    img.save(p)
    box = (35, 35, 110, 50)      # 比文字略大，环以填充色为主
    default = hl_layout.measure_coverage(p, box)
    inflated = hl_layout.measure_coverage(p, box, bg_rgb=(255, 255, 255))
    assert 0.3 < default < 0.6, f"默认口径应只算文字，实测 {default:.3f}"
    assert inflated > 0.9, f"主色口径应把整块填充算成墨迹，实测 {inflated:.3f}"
    assert inflated > default + 0.3
    # 独立口径（环取 rect 外的填充色）与契约默认一致
    assert _far_ring_coverage(p, box, far=4) == pytest.approx(default, abs=0.08)


def test_ring_bg_estimator_tilts_when_ink_covers_rect_edges(tmp_path):
    """把上一条推到极端：文字**铺满 rect 的上下边**时，内侧环的中位色倒向墨迹色，

    默认口径从 0.6 掉到 0.4（环上 200/280 像素是墨迹）。本素材实测环上墨迹占比
    ≤0.037（见 tools/probes/indep_m2_coverage.py），所以真稿上不发作 —— 但"墨迹
    触框边"（impl 报告 12/22 行）离这个失效模式并不远，验收脚本必须显式传局部底色。
    """
    from PIL import Image, ImageDraw
    p = str(tmp_path / "edge_ink.png")
    img = Image.new("RGB", (200, 120), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.rectangle([20, 20, 179, 99], fill=(182, 199, 220))
    for y in (40, 56, 72):
        d.rectangle([40, y, 139, y + 7], fill=(20, 20, 20))
    img.save(p)
    tight = (40, 40, 100, 40)                     # rect 恰好贴合文字块
    bars = 3 * 8 * 100                            # 三行 8px 高、100px 宽的文字
    truth = bars / (tight[2] * tight[3])           # 真实覆盖率 = 0.60
    got = hl_layout.measure_coverage(p, tight)
    assert got < truth - 0.15, f"环被墨迹铺满时应显著低估：测得 {got:.3f} vs 真值 {truth:.3f}"
    # 而"环以填充色为主"的稍大 rect 测得接近真值（0.436 ≈ 2400/5500）
    loose = (35, 35, 110, 50)
    assert hl_layout.measure_coverage(p, loose) == pytest.approx(2400 / (110 * 50), abs=0.06)


@pytest.mark.skipif(not os.path.isfile(B_MULTISLIDE), reason="缺 output/b_multislide.pptx")
def test_real_material_median_coverage_reproduced_independently():
    """真实素材：独立口径复算行级覆盖率中位 ≈ 0.286（容差 0.04）。

    同时确认"上限 = 墨迹 bbox 内密度中位 ≈ 0.414"→ 契约阈值 0.35 用矩形法不可达。
    """
    import statistics
    from PIL import Image
    deck, _ = pptx_io.read_pages(B_MULTISLIDE)
    covs, dens = [], []
    for page in deck.pages:
        bg = os.path.join(BG_DIR, f"slide_{page['index'] + 1}.png")
        if not os.path.isfile(bg):
            pytest.skip(f"缺底图 {bg}（先跑 tools/probes/accept_m1.py）")
        img = Image.open(bg).convert("RGB")
        for u in hl_layout.build_units(hl_layout.page_shapes(page, deck)):
            for r in u.lines:
                box = (int(r.left_px), int(r.top_px), int(r.width_px), int(r.height_px))
                if box[2] <= 0 or box[3] <= 0:
                    continue
                stats = _bbox_and_ink(img, box)
                if stats is None:
                    continue
                ink, (bw, bh) = stats
                covs.append(ink / (box[2] * box[3]))
                dens.append(ink / (bw * bh))
    assert covs, "没有测到任何行 rect"
    med = statistics.median(covs)
    assert med == pytest.approx(0.286, abs=0.04), f"独立复算 coverage 中位 {med:.3f}"
    assert statistics.median(dens) < 0.45, "墨迹 bbox 密度上限应远低于 0.5"
    assert med < 0.35, "本素材 coverage 中位确实达不到契约阈值 0.35"


def _bbox_and_ink(img, box, thresh=40, far=8):
    """返回 (rect 内墨迹像素数, 墨迹 tight bbox 尺寸)；无墨迹返回 None。"""
    W, H = img.size
    l, t, w, h = box
    r, b = l + w, t + h
    px = img.load()
    coords = []
    for band in range(far, far + 3):
        for x in range(max(0, l - band), min(W, r + band)):
            for y in (t - band, b + band - 1):
                if 0 <= y < H:
                    coords.append((x, y))
        for y in range(max(0, t - band), min(H, b + band)):
            for x in (l - band, r + band - 1):
                if 0 <= x < W:
                    coords.append((x, y))
    counts = {}
    for x, y in coords:
        p = px[x, y]
        k = (p[0] // 8, p[1] // 8, p[2] // 8)
        counts[k] = counts.get(k, 0) + 1
    if not counts:
        return None
    k = max(counts, key=counts.get)
    bg = (k[0] * 8 + 4, k[1] * 8 + 4, k[2] * 8 + 4)
    ink, x0, x1, y0, y1 = 0, 10 ** 9, -1, 10 ** 9, -1
    for y in range(max(0, t), min(H, b)):
        for x in range(max(0, l), min(W, r)):
            p = px[x, y]
            dr, dg, db = p[0] - bg[0], p[1] - bg[1], p[2] - bg[2]
            if dr * dr + dg * dg + db * db > thresh * thresh:
                ink += 1
                x0, x1, y0, y1 = min(x0, x), max(x1, x), min(y0, y), max(y1, y)
    if x1 < 0:
        return None
    return ink, (x1 - x0 + 1, y1 - y0 + 1)


# ================================================================ 2. 断行等价

CORPUS = [
    "a  b   c     中文  结尾",              # 连续空格
    "。，、！？；：「测试」全角标点",          # 全角标点行首
    "😀😀😀abc测试🎉",                      # emoji
    "a\tb\tc\t中文",                        # 制表符
    "第一行\r\n第二行\r\n第三行",             # CRLF
    "1234567890" * 8,                       # 纯数字长串
    "supercalifragilisticexpialidocious中文混排",  # 超长英文词 + 中文
    "中文　　全角空格　　结束",               # U+3000
    "a​b‌c",                      # 零宽字符
    "café näive",               # 组合音标
    "——————" * 6,                          # 长破折号
    "----_____++++====",                    # ASCII 标点串
    "https://example.com/a/very/long/path?c=d&e=f#frag",
    "\n\nabc",                              # 换行前缀
    "abc\n\n",                              # 换行后缀
    "     ",                                # 只有空格
    "",                                     # 空串
    "测试 한국어 テスト",                     # 日韩
]


@pytest.mark.parametrize("text", CORPUS)
@pytest.mark.parametrize("size", [8.0, 12.0, 18.0, 40.0])
@pytest.mark.parametrize("width", [30.0, 90.0, 250.0, 800.0])
def test_wrap_lines_matches_qa_at(text, size, width):
    """等价性（独立语料）：是我挑的、实现者测试里没有的文本。"""
    assert len(hl_layout.wrap_lines(text, size, width)) == \
        qa.measure_text_lines(text, size, width)


@pytest.mark.parametrize("text", CORPUS)
@pytest.mark.parametrize("width", [45.0, 150.0, 600.0])
def test_wrap_lines_loses_no_non_space_char(text, width):
    """不变量：除空格外的字符一个都不能丢（抓"掉字"类真 bug）。"""
    lines = hl_layout.wrap_lines(text, 14.0, width)
    assert "".join(ln.text for ln in lines).replace(" ", "") == text.replace(" ", "")


@pytest.mark.parametrize("text", CORPUS)
@pytest.mark.parametrize("width", [45.0, 150.0, 600.0])
def test_wrap_lines_no_leading_space_and_no_empty_line(text, width):
    lines = hl_layout.wrap_lines(text, 14.0, width)
    assert not any(ln.text.startswith(" ") for ln in lines)
    if text.strip(" "):
        assert all(ln.text for ln in lines)


@pytest.mark.parametrize("text", ["疆", "W", "supercalifragilistic", "中文测试"])
def test_wrap_lines_only_single_char_lines_may_exceed_width(text):
    """不变量：行宽超上限只允许发生在"单个字符本身超宽"的强拆行。"""
    width = 20.0
    for ln in hl_layout.wrap_lines(text, 40.0, width):
        if ln.width_pt > width + 1e-9:
            assert len(ln.text) == 1


def test_word_wrap_false_is_intentional_divergence():
    """契约 §4.3：word_wrap=False 不换行 → 与 qa（永远按换行算）**有意分歧**。"""
    long_text = "很长的一句话。" * 30
    assert len(hl_layout._paragraph_lines(long_text, 18.0, 100.0, False)) == 1
    assert qa.measure_text_lines(long_text, 18.0, 100.0) > 1
    assert len(hl_layout._paragraph_lines(long_text, 18.0, 100.0, True)) == \
        qa.measure_text_lines(long_text, 18.0, 100.0)


# ================================================================ 3. 转义对抗

def _unit(text="t", rect=None):
    r = rect or Rect(1270000, 635000, 2540000, 508000, 200.0, 100.0, 400.0, 80.0)
    return Unit(order=0, page_index=0, text=text, kind="body", rect=r, lines=[r],
                size_pt=18.0, align="LEFT", shape_id=1, shape_name="S",
                is_estimated=False, warnings=[])


TITLE_PAYLOAD = "</title><script>alert(1)</script>__CONFIG__<!--x-->"


def test_adversarial_title_and_units_static(tmp_path):
    pages = [[_unit("</script><script>alert(2)</script>"),
              _unit("<!-- comment -->"),
              _unit("</SCRIPT>alert(3)</SCRIPT>"),
              _unit("<img src=x onerror=alert(4)>"),
              _unit("__TITLE__ __BGJSON__ __UNITS__ __CONFIG__"),
              _unit("  行分隔   段分隔"),
              _unit('"}; alert(5); var x={"'),
              _unit("\x00 空字节")]]
    path = hl_anim.build_player(str(tmp_path / "p"), ["bg/slide_1.png"], pages,
                                title=TITLE_PAYLOAD)
    doc = open(path, encoding="utf-8").read()

    assert doc.lower().count("</script") == 1, "脚本块被提前闭合"
    assert "<iframe" not in doc and "sandbox" not in doc
    head, _, rest = doc.partition("<script>")
    body, _, tail = rest.rpartition("</script>")
    assert "<!--" not in body, "`<!--` 未转义 → 可进 double-escaped 态"
    for bad in ("onerror=", "onload="):
        assert bad not in (head + tail).lower()

    def to_json(s):
        return s.replace("<\\/", "</").replace("<\\!--", "<!--")

    import json
    for pat in (r"const CFG = (.*?);\n", r"const BGS = (.*?);\n", r"const PAGES = (.*?);\n"):
        m = re.search(pat, doc)
        assert m, f"找不到 {pat}"
        json.loads(to_json(m.group(1)))          # 还原 JS 转义后必须仍是合法 JSON
    assert "__CONFIG__" in doc, "title 里的占位符字面量应原样保留（单遍替换）"


def test_adversarial_bg_paths_encoded_and_traversal_blocked(tmp_path):
    from pptx_io import PptxError
    import json
    path = hl_anim.build_player(str(tmp_path / "p2"),
                                ['bg/a"><script>x</script>.png'], [[]])
    doc = open(path, encoding="utf-8").read()
    assert "<script>x</script>" not in doc
    encoded = json.loads(re.search(r"const BGS = (.*?);\n", doc).group(1))[0]
    assert not (set('<>"# ') & set(encoded))
    assert re.fullmatch(r"([^%]|%[0-9A-Fa-f]{2})+", encoded), "裸 % 会被当转义引导符"

    for bad in ("..\\..\\x.png", "bg\\..\\..\\x.png", "C:foo.png",
                "\\\\server\\share\\x.png", "//server/share/x.png"):
        with pytest.raises(PptxError) as ei:
            hl_anim.build_player(str(tmp_path / "p3"), [bad], [[]])
        assert ei.value.code == "IR_MISMATCH"


def _browser_or_skip(p):
    import shot
    try:
        return shot._launch_browser(p)
    except RuntimeError as exc:
        pytest.skip(f"无可用浏览器：{exc}")


def test_adversarial_player_in_real_browser(tmp_path):
    """真实 Chrome：恶意 title 按字面渲染、无 pageerror、无对话框、goto 同步。"""
    from playwright.sync_api import sync_playwright
    from PIL import Image

    root = tmp_path / "player"
    (root / "bg").mkdir(parents=True)
    hostile = "slide_#1%x 中文.png"
    Image.new("RGB", (64, 36), (12, 34, 56)).save(root / "bg" / hostile)

    path = hl_anim.build_player(str(root), [f"bg/{hostile}", "bg/missing.png"],
                                [[_unit("正文")], [_unit("第二页")]],
                                title=TITLE_PAYLOAD)
    errors, dialogs = [], []
    with sync_playwright() as p:
        browser = _browser_or_skip(p)
        try:
            page = browser.new_page(viewport={"width": 1280, "height": 720})
            page.on("pageerror", lambda e: errors.append(str(e)))
            page.on("dialog", lambda d: (dialogs.append(d.message), d.dismiss()))
            page.goto("file:///" + path.replace("\\", "/"), wait_until="load")
            page.evaluate("window.hl.ready")
            assert page.evaluate("document.title") == TITLE_PAYLOAD + " · 高亮讲解"
            assert page.evaluate("document.querySelector('h1').textContent") == \
                TITLE_PAYLOAD + " · 高亮讲解"
            page.evaluate("window.hl.goto(0, 1)")
            assert page.evaluate("window.hl.state()")["step"] == 1   # goto 同步生效
            assert page.evaluate(
                "Array.from(document.querySelectorAll('.hl'))"
                ".filter(d=>d.classList.contains('on')).length") == 1
            loaded = page.evaluate(
                "Array.from(document.querySelectorAll('img.bg')).map(i=>i.naturalWidth)")
        finally:
            browser.close()
    assert not errors, f"页面 JS 报错：{errors}"
    assert not dialogs, f"注入被执行（弹出对话框）：{dialogs}"

    assert loaded[0] > 0, "含 #/%/空格/中文 的底图路径未取到真实文件（编码有问题）"


# ================================================================ 4. 病态输入

def test_pathological_inputs_give_chinese_pptxerror(tmp_path):
    from pptx_io import PptxError
    import zipfile as zf

    cases = []
    cases.append((str(tmp_path / "nope.pptx"), "PPTX_NOT_FOUND"))
    txt = tmp_path / "renamed.pptx"
    txt.write_text("不是 pptx", encoding="utf-8")
    cases.append((str(txt), "PPTX_UNREADABLE"))
    ole = tmp_path / "ole.pptx"
    # 真加密的 OOXML 是 OLE2 + `EncryptedPackage` 特征流（CFB 目录以 UTF-16LE 存名）。
    # 只写一个 OLE2 头是**假样本** —— 那等于断言"任何 OLE2 都算加密"，正是审计 L1
    # 指出的误判（老 .ppt 也是 OLE2）。这里补上特征流，样本才代表真实加密稿。
    ole.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(64)
                    + "EncryptedPackage".encode("utf-16-le") + bytes(64))
    cases.append((str(ole), "PPTX_ENCRYPTED"))
    trunc = tmp_path / "trunc.pptx"
    trunc.write_bytes(b"PK\x03\x04" + bytes(100))
    cases.append((str(trunc), "PPTX_UNREADABLE"))
    notzip = tmp_path / "notapptx.pptx"
    with zf.ZipFile(notzip, "w") as z:
        z.writestr("hello.txt", "hi")
    cases.append((str(notzip), "PPTX_UNREADABLE"))

    for path, code in cases:
        with pytest.raises(PptxError) as ei:
            pptx_io.read_pages(path)
        assert ei.value.code == code, f"{path} → {ei.value.code}"
        assert ei.value.message and any("一" <= c <= "鿿" for c in ei.value.message)
        assert all(ord(c) < 128 for c in ei.value.code)


def test_zero_slide_pptx_is_pptx_empty(tmp_path):
    from pptx_io import PptxError
    p = str(tmp_path / "zero.pptx")
    Presentation().save(p)
    with pytest.raises(PptxError) as ei:
        pptx_io.read_pages(p)
    assert ei.value.code == "PPTX_EMPTY"


def test_slide_without_shapes_yields_no_units(tmp_path):
    p = str(tmp_path / "oneblank.pptx")
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(p)
    deck, _ = pptx_io.read_pages(p)
    assert len(deck.pages) == 1
    assert hl_layout.build_units(hl_layout.page_shapes(deck.pages[0], deck)) == []


def test_extreme_content_does_not_crash(tmp_path):
    """零尺寸 / 隐藏 / 纯空白 / 超宽单字 / word_wrap=False：都不崩，且按契约过滤。"""
    p = str(tmp_path / "edges.pptx")
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "     "
    tb2 = s.shapes.add_textbox(Inches(1), Inches(2), Inches(1), Inches(1))
    tb2.name = "HugeChar"
    r = tb2.text_frame.paragraphs[0].add_run()
    r.text, r.font.size = "疆", Pt(200)
    z = s.shapes.add_textbox(Inches(9), Inches(1), Emu(0), Inches(1))
    z.name, z.text_frame.text = "ZeroWidth", "零宽框"
    h = s.shapes.add_textbox(Inches(9), Inches(3), Inches(2), Inches(1))
    h.name, h.text_frame.text = "HiddenBox", "隐藏文本"
    h._element.nvSpPr.cNvPr.set("hidden", "1")
    prs.save(p)

    deck, skips = pptx_io.read_pages(p)
    reasons = {s["reason"] for s in skips}
    assert {"empty_text", "zero_size", "hidden"} <= reasons
    units = hl_layout.build_units(hl_layout.page_shapes(deck.pages[0], deck))
    huge = [u for u in units if u.shape_name == "HugeChar"]
    assert huge and len(huge[0].lines) == 1        # 超宽单字：强拆但只占 1 行


def test_no_powerpoint_gives_clear_chinese_error(tmp_path, monkeypatch):
    """打桩：无 PowerPoint 时 export_pages 给明确中文 PptxError，不是裸堆栈。"""
    from pptx_io import PptxError
    monkeypatch.setattr(pptx_io, "powerpoint_available", lambda: False)
    with pytest.raises(PptxError) as ei:
        pptx_io.export_pages(B_MULTISLIDE, str(tmp_path / "bg"))
    assert ei.value.code == "NO_POWERPOINT"
    assert "PowerPoint" in ei.value.message
    assert ei.value.hint, "必须带可执行指引"


# ================================================================ 5. M1 口径

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


@pytest.mark.skipif(not os.path.isfile(B_MULTISLIDE), reason="缺 output/b_multislide.pptx")
def test_m1_shape_counts_verified_by_raw_xml():
    """裸 XML 数形状：41 = 40 文本框 + 1 图表；过滤空文本后保留 23（22 text + 1 chart）。"""
    from lxml import etree
    tot = Counter()
    with zipfile.ZipFile(B_MULTISLIDE) as z:
        names = [n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        for n in names:
            tree = etree.fromstring(z.read(n)).find(f"{P}cSld/{P}spTree")
            for el in tree:
                tag = etree.QName(el).localname
                if tag == "sp":
                    tot["sp"] += 1
                elif tag == "graphicFrame" and el.find(f".//{C}chart") is not None:
                    tot["chart"] += 1
    assert tot["sp"] == 40 and tot["chart"] == 1 and sum(tot.values()) == 41

    deck, skipped = pptx_io.read_pages(B_MULTISLIDE)
    kinds = Counter(s.kind for pg in deck.pages for s in pg["shapes"])
    assert len(deck.pages) == 10
    assert dict(kinds) == {"text": 22, "chart": 1}
    assert Counter(s["reason"] for s in skipped) == {"empty_text": 18}
    assert sum(kinds.values()) + len(skipped) == 41
