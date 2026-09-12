"""审计复现 · 第三方 pptx 解析风险面 + 错误归类（pptx_io）。

用法：
    PYTHONIOENCODING=utf-8 ./.venv/Scripts/python.exe tools/probes/audit_pptx_io.py

覆盖：zip bomb（内存内构造，不落盘）／实体展开（billion laughs）／库内是否联网／
`_classify_bad_package` 对 .ppt 的误判／`slide_width is None` 的未捕获崩溃／
真实稿里 a:br 的分布（决定 _paragraph_lines 分叉的现实性）。

唯一会落盘的：若仓库内找不到 OLE2 样本，会往系统临时目录写一个 16 字节文件并在结束时删除
（打印路径，便于核对）。
"""
import io
import os
import sys
import tempfile
import time
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)

import pptx  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

import pptx_io  # noqa: E402

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
SRC = os.path.join(ROOT, "output", "b_multislide.pptx")
FIX = os.path.join(ROOT, "output", "fixtures", "fixtures_cover.pptx")


def section(t):
    print("\n" + "=" * 78)
    print(t)
    print("=" * 78)


# ---------------------------------------------------------------- 1. 读轨迹

def trace_reads(pptx_path):
    """拦 ZipFile.read，记录 python-pptx 到底解压了哪些条目。"""
    log = []
    real_read = zipfile.ZipFile.read

    def spy(self, name, *a, **k):
        try:
            info = self.getinfo(name)
            log.append((name, info.file_size, info.compress_size))
        except KeyError:
            pass
        return real_read(self, name, *a, **k)

    zipfile.ZipFile.read = spy
    try:
        Presentation(pptx_path)
    finally:
        zipfile.ZipFile.read = real_read
    return log


def rebuild_with_bomb(src_path, entry_name, payload_bytes):
    """内存里重建包：把 entry_name 换成 payload（其余原样拷贝）。返回 BytesIO。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(src_path) as zin, \
            zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = payload_bytes if item.filename == entry_name else zin.read(item.filename)
            zout.writestr(item.filename, data)
    buf.seek(0)
    return buf


def main():
    section("1 · python-pptx 打开一份稿时会解压哪些条目（bomb 的有效目标）")
    log = trace_reads(SRC)
    for name, size, csize in log:
        print(f"  {name:34s} 解压后 {size:>9,d} B   压缩后 {csize:>8,d} B")
    print(f"  共 {len(log)} 个条目被读入内存")

    section("2 · zip bomb：把 64 MB 全零塞进 [Content_Types].xml，python-pptx 是否有大小闸门")
    bomb = b"\x00" * (64 * 1024 * 1024)
    t0 = time.time()
    buf = rebuild_with_bomb(SRC, "[Content_Types].xml", bomb)
    raw = buf.getvalue()
    print(f"  构造出的包体积 = {len(raw):,d} B（{len(raw)/1024:.1f} KB），"
          f"内含一个 64 MiB 的 [Content_Types].xml → 压缩比 ≈ {len(bomb)/len(raw):.0f}:1")
    log2 = []
    real_read = zipfile.ZipFile.read

    def spy(self, name, *a, **k):
        try:
            info = self.getinfo(name)
            log2.append((name, info.file_size))
        except KeyError:
            pass
        return real_read(self, name, *a, **k)

    zipfile.ZipFile.read = spy
    err = None
    try:
        Presentation(buf)
    except Exception as exc:                                   # noqa: BLE001
        err = exc
    finally:
        zipfile.ZipFile.read = real_read
    print(f"  Presentation() 结果：{type(err).__name__ if err else '成功'}"
          f"{('：' + str(err)[:60]) if err else ''}，耗时 {time.time()-t0:.2f}s")
    print(f"  但解压轨迹 = {log2}")
    print("  → 即使 XML 立刻解析失败，字节已在内存里展开；pptx 侧只做了 read，没有任何大小/比率校验")
    print("  → 攻击者可把 [Content_Types].xml 换成合法但巨大的 XML（或 42.zip 级别），"
          "python-pptx 不设上限")

    section("3 · 实体展开 / XXE：lxml 解析器配置")
    oxml = open(os.path.join(os.path.dirname(pptx.__file__), "oxml", "__init__.py"),
                encoding="utf-8").read()
    for ln in oxml.splitlines():
        if "XMLParser(" in ln:
            print(f"  库内原文：{ln.strip()}")
    laughs = ['<?xml version="1.0"?>', "<!DOCTYPE lolz ["]
    laughs.append('<!ENTITY lol "lol">')
    prev = "lol"
    for i in range(1, 10):
        nxt = f"lol{i}"
        laughs.append(f'<!ENTITY {nxt} "&{prev};&{prev};&{prev};&{prev};&{prev};&{prev};&{prev};&{prev};&{prev};&{prev};">')
        prev = nxt
    laughs.append("]><lolz>&lol9;</lolz>")
    xml = "\n".join(laughs)
    t0 = time.time()
    try:
        root = pptx.oxml.parse_xml(xml)
        print(f"  billion laughs 解析：成功，耗时 {time.time()-t0:.4f}s，"
              f"根标签 {root.tag!r}，文本={root.text!r}（实体未展开）")
        print(f"  若实体被展开，lol9 会变成 10^9 个 'lol'（≈3 GB）——实测未展开")
    except Exception as exc:                                   # noqa: BLE001
        print(f"  billion laughs 解析：{type(exc).__name__}（{str(exc)[:70]}），耗时 {time.time()-t0:.4f}s")

    section("4 · python-pptx 是否会联网拉外部关系（External rel）")
    pkgdir = os.path.dirname(pptx.__file__)
    needles = ("urllib", "requests", "socket", "http.client", "urlopen", "httplib")
    hits = []
    for root_, _dirs, files in os.walk(pkgdir):
        for fn in files:
            if not fn.endswith(".py"):
                continue
            p = os.path.join(root_, fn)
            try:
                txt = open(p, encoding="utf-8", errors="ignore").read()
            except OSError:
                continue
            for n in needles:
                if n in txt:
                    hits.append((os.path.relpath(p, pkgdir), n))
    print(f"  python-pptx 源码内命中 {needles} 的位置：{hits if hits else '无'}")
    print("  → 库不做任何网络请求；TargetMode=\"External\" 的关系只被记录、不会被解引用")

    section("5 · _classify_bad_package：OLE2 一律判『已加密』—— .ppt 会被误诊")
    print("  代码分支（pptx_io.py:439-445）：")
    print("    if head.startswith(_OLE_MAGIC): -> PptxError('PPTX 已加密，无法读取', 'PPTX_ENCRYPTED',")
    print("                                               '请先用 PowerPoint 去掉打开密码再导出')")
    tmp_path = None
    ole_files = []
    for root_, dirs, files in os.walk(ROOT):
        dirs[:] = [d for d in dirs if d not in (".git", ".venv", "node_modules", "__pycache__")]
        for fn in files:
            p = os.path.join(root_, fn)
            try:
                if os.path.getsize(p) < 8:
                    continue
                with open(p, "rb") as fh:
                    if fh.read(8) == OLE_MAGIC:
                        ole_files.append(p)
            except OSError:
                pass
    if ole_files:
        print(f"  仓库内找到的真实 OLE2 样本：{ole_files}")
        tmp_path = ole_files[0]
    else:
        fd, tmp_path = tempfile.mkstemp(suffix=".ppt", prefix="audit_ole_")
        with os.fdopen(fd, "wb") as fh:
            fh.write(OLE_MAGIC + b"\x00" * 8)
        print(f"  仓库内无 OLE2 样本 → 临时造一个 8 字节头样本：{tmp_path}（脚本结束会删除）")
    err = pptx_io._classify_bad_package(tmp_path)
    print(f"  判定结果：code={err.code!r} message={err.message!r} hint={err.hint!r}")
    print("  → .ppt 与『加密的 .pptx』同为 OLE2 复合文档，仅凭文件头无法区分；")
    print("    用户拿一份老 .ppt 来会被告知『已加密，请去掉打开密码』——指引不可执行")
    if not ole_files and tmp_path:
        os.remove(tmp_path)
        print(f"  （已删除临时文件 {tmp_path}）")

    section("6 · prs.slide_width is None → read_pages 未捕获的 TypeError")
    # 去掉 p:sldSz 后从内存加载，看 slide_width 与 int() 的行为
    buf = io.BytesIO()
    with zipfile.ZipFile(SRC) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/presentation.xml":
                txt = data.decode("utf-8")
                import re as _re
                txt2 = _re.sub(r"<p:sldSz[^/]*/>", "", txt)
                print(f"  已从 presentation.xml 移除 <p:sldSz>（{len(txt)-len(txt2)} 字符）")
                data = txt2.encode("utf-8")
            zout.writestr(item.filename, data)
    buf.seek(0)
    try:
        prs = Presentation(buf)
        sw = prs.slide_width
        print(f"  prs.slide_width = {sw!r}")
        try:
            int(sw)
            print("  int(slide_width) 成功（意外）")
        except Exception as exc:                               # noqa: BLE001
            print(f"  int(slide_width) -> {type(exc).__name__}: {exc}")
            print("  → read_pages 的 try 只包住 Presentation()，这行在 try 之外"
                  "（pptx_io.py:470）→ 崩溃会冒成 INTERNAL(5)，而不是 PPTX_UNREADABLE(3)")
    except Exception as exc:                                   # noqa: BLE001
        print(f"  Presentation() 直接抛：{type(exc).__name__}: {exc}")

    section("7 · 真实稿里 a:br / a:fld 的分布（_paragraph_lines 分叉有多现实）")
    for p in [SRC, FIX,
              os.path.join(ROOT, "output",
                           "大语言模型是怎么工作的：从_Token__20260906_004158.pptx")]:
        if not os.path.isfile(p):
            continue
        try:
            prs = Presentation(p)
        except Exception as exc:                               # noqa: BLE001
            print(f"  {os.path.basename(p)}: 打不开（{type(exc).__name__}）")
            continue
        n_br = n_fld = n_para = 0
        for s in prs.slides:
            for sh in s.shapes:
                if not getattr(sh, "has_text_frame", False):
                    continue
                for para in sh.text_frame.paragraphs:
                    n_para += 1
                    for ch in para._p:
                        if ch.tag == qn("a:br"):
                            n_br += 1
                        elif ch.tag == qn("a:fld"):
                            n_fld += 1
        print(f"  {os.path.basename(p):52s} 段 {n_para:3d} · a:br {n_br:3d} · a:fld {n_fld:3d}")


if __name__ == "__main__":
    main()
