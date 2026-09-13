"""M4 对照探针：缺 `p:sldSz` 的包会让错误码从 3 退化成 5。

背景（AUDIT_REPORT §2 M4）：`prs.slide_width` 可以是 `None`（包里没有 `<p:sldSz>`），
而 `int(prs.slide_width)` 那行在 `read_pages` 的 try **之外** → 裸 `TypeError` 冒出去，
最终被 CLI 归成 `INTERNAL(5)`，把"用户稿子畸形"说成"这是我们的 bug"。

本探针纯内存/临时文件操作，**不启动任何外部程序**（无需 COM 安全闸门）。

判定：
  修复前 → 抛 TypeError（不是 PptxError）→ CLI 只能归 INTERNAL(5)
  修复后 → 抛 PptxError，code=PPTX_UNREADABLE（退出码 3）、message 说明缺 sldSz

跑法：.venv/Scripts/python.exe tools/probes/fix_m4_sldsz.py
"""

import os
import re
import shutil
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, ROOT)

SLD_SZ_RE = re.compile(r"<p:sldSz\b[^>]*/>")


def make_deck_without_sldsz(path: str) -> int:
    """造一份真 pptx，再把 presentation.xml 里的 <p:sldSz/> 删掉。"""
    from pptx import Presentation
    src = path + ".src.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(src)

    removed = 0
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/presentation.xml":
                text = data.decode("utf-8")
                text, removed = SLD_SZ_RE.subn("", text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    os.remove(src)
    return removed


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="m4probe_")
    try:
        deck = os.path.join(tmp, "no_sldsz.pptx")
        removed = make_deck_without_sldsz(deck)
        print(f"样本：{deck}")
        print(f"  已从 presentation.xml 移除 <p:sldSz>：{removed} 处"
              f"（{'成功' if removed else '**没删到，样本无效**'}）")
        if not removed:
            return 1

        from pptx import Presentation
        print(f"  自检：prs.slide_width = {Presentation(deck).slide_width!r}（应为 None）")

        import app as app_mod  # noqa: F401  仅为确认 cli 的码表可达
        import cli
        import pptx_io

        print("\n== 直接调 read_pages ==")
        code = None
        try:
            pptx_io.read_pages(deck)
            print("  结果：**没报错**？不应发生")
        except pptx_io.PptxError as exc:
            code = exc.code
            print(f"  结果：PptxError code={exc.code}")
            print(f"        message={exc.message}")
            print(f"        hint={exc.hint}")
        except Exception as exc:  # noqa: BLE001
            print(f"  结果：**裸异常** {type(exc).__name__}: {exc}")
            print("        → CLI 只能归 INTERNAL(5)，把用户稿畸形说成我们的 bug")

        print("\n== 走 CLI（看最终退出码）==")
        rc = cli.main(["import", deck, "--no-com", "--mode", "redesign",
                       "--out", os.path.join(tmp, "out")])
        print(f"  CLI 退出码 = {rc}（修复前 5=INTERNAL；修复后应为 3=输入问题）")

        print("\n" + "-" * 66)
        if code == "PPTX_UNREADABLE" and rc == 3:
            print("判定：**FIXED** —— 错误码 PPTX_UNREADABLE(3)，不再冒充内部错误")
        else:
            print("判定：**REPRODUCED** —— 错误码退化成 INTERNAL(5)")
        print("-" * 66)
        return 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
